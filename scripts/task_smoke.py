#!/usr/bin/env python
# -*- coding: utf-8 -*-
"""任务层冒烟（TASK-002D）—— 起**真服务**、发**真 HTTP**、验**真文件**。

为什么要单独一个脚本（而不是并进 api_smoke.py）：
    api_smoke.py 是 002C 的验收脚本（6 个端点，Hermes 也用它跑独立冒烟），
    改它等于动别人的验收基线 —— 所以这里**照它的骨架另立一份**，专测 002D 的任务链路。
    pytest 用的是 `TestClient`（进程内直调），**证明不了"真的走 HTTP"**；
    本脚本起真 uvicorn、发真请求、真下载并用 openpyxl 读回数字，作为 002D 的 GATE 证据。

跑法：
    .venv/Scripts/python.exe scripts/task_smoke.py            # 默认端口 8501
    .venv/Scripts/python.exe scripts/task_smoke.py --port 8510
预期末行：`🎉 任务层冒烟全部通过 …`
"""

from __future__ import annotations

import argparse
import hashlib
import json
import os
import re
import subprocess
import sys
import tempfile
import time

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
os.chdir(ROOT)
ENV = dict(os.environ)
ENV["PYTHONPATH"] = ""
PY = os.path.join(ROOT, ".venv", "Scripts", "python.exe")

import httpx  # noqa: E402

START, END = "2011-11-21", "2011-11-27"
METRICS = ["sales_amount", "rows_in_range", "rows_valid", "rows_excluded"]
EXPECTED_ROWS_IN_RANGE = 19950
EXPECTED_ROWS_EXCLUDED = 296
EXPECTED_AMOUNT = 316412.16

# 服务器绝对路径不得出现在对外响应里（R005 required_change #7）。
# 前面那个 (?<![A-Za-z]) 是排除 "https://" 里的 "s:/"（否则每个带 scheme 的 URL 都误报）。
ABS_PATH_RE = re.compile(r"(?<![A-Za-z])[A-Za-z]:[\\/]{1,2}")

failures: list[str] = []


def check(condition: bool, label: str, extra: str = "") -> None:
    """断言但**不中止**：一条失败也继续跑完，最后统一汇总（排障时不用反复重跑几分钟的任务链）。"""
    if condition:
        print(f"   ✅ {label}{(' | ' + extra) if extra else ''}")
    else:
        print(f"   ❌ {label}{(' | ' + extra) if extra else ''}")
        failures.append(label)


def main() -> int:
    parser = argparse.ArgumentParser(description="任务层真服务冒烟（TASK-002D）")
    parser.add_argument("--port", type=int, default=8501)
    args = parser.parse_args()
    base = f"http://127.0.0.1:{args.port}"

    # **隔离落盘目录**（state.py 的 SRA_* 环境变量就是为这个准备的）：
    # 冒烟会故意建一个"非快照数据"的任务并跑出一次失败记录 —— 这种垃圾记录不该进
    # state/*.json（那是真审计凭据）。跑在临时目录里，真状态一条不脏，且目录路径会打印出来供人查看。
    sandbox = tempfile.mkdtemp(prefix="sra_task_smoke_", dir=os.path.join(ROOT, "outputs"))
    ENV["SRA_STATE_DIR"] = os.path.join(sandbox, "state")
    ENV["SRA_UPLOAD_DIR"] = os.path.join(sandbox, "uploads")
    ENV["SRA_OUTPUT_DIR"] = os.path.join(sandbox, "outputs")
    os.makedirs(ENV["SRA_STATE_DIR"], exist_ok=True)
    os.makedirs(ENV["SRA_UPLOAD_DIR"], exist_ok=True)
    os.makedirs(ENV["SRA_OUTPUT_DIR"], exist_ok=True)
    print(f"落盘沙箱（跑完可删）：{sandbox}")

    proc = subprocess.Popen(
        [PY, "-m", "uvicorn", "app.api:app", "--host", "127.0.0.1", "--port", str(args.port)],
        stdout=subprocess.PIPE, stderr=subprocess.STDOUT, env=ENV,
    )
    try:
        print("=== 0) 等 service 就绪 ===")
        ready = False
        for _ in range(60):
            time.sleep(2)
            try:
                response = httpx.get(f"{base}/api/health", timeout=3)
                if response.status_code == 200:
                    health = response.json()
                    print(f"   HTTP 200 | version={health.get('version')} | task={health.get('task')}")
                    check(health.get("version") == "0.2.0", "版本是 0.2.0（002D）", str(health.get("version")))
                    ready = True
                    break
            except Exception:
                pass
        if not ready:
            print("❌ service 未就绪，放弃")
            return 1

        print("\n=== 1) 真上传 23MB xlsx ===")
        started = time.time()
        with open(os.path.join(ROOT, "data", "Online Retail.xlsx"), "rb") as handle:
            response = httpx.post(f"{base}/api/upload", files={"file": ("Online Retail.xlsx", handle)}, timeout=900)
        upload = response.json()
        print(f"   HTTP {response.status_code} | file_id={upload.get('file_id')} | rows={upload.get('rows')} | "
              f"{time.time() - started:.1f}s")
        check(response.status_code == 201, "上传 HTTP 201", str(response.status_code))
        check(upload.get("rows") == 541909, "上传行数 541909", str(upload.get("rows")))
        file_id = upload.get("file_id")

        print("\n=== 2) 建任务（便捷路径：file_id + 区间 + 指标）===")
        started = time.time()
        response = httpx.post(
            f"{base}/api/tasks",
            json={"file_id": file_id, "name": "冒烟·上周销售汇总", "start": START, "end": END, "metrics": METRICS},
            timeout=300,
        )
        task = response.json()
        print(f"   HTTP {response.status_code} | task_id={task.get('task_id')} | status={task.get('status')} | "
              f"spec_version={task.get('spec_version')} | {time.time() - started:.1f}s")
        check(response.status_code == 201, "建任务 HTTP 201", str(response.status_code))
        task_id = task.get("task_id")
        check(bool(task_id), "返回 task_id", str(task_id))
        check(task.get("status") == "created", "初始状态 created", str(task.get("status")))
        check("warnings" not in task, "数据与快照一致 → 无 warnings")
        spec = task.get("spec") or {}
        check(spec.get("time_range", {}).get("mode") == "absolute", "冻结的 Spec 带 time_range.mode=absolute")
        check(spec.get("time_range", {}).get("start") == START and spec.get("time_range", {}).get("end") == END,
              "区间写进了 Spec", f"{spec.get('time_range', {}).get('start')}~{spec.get('time_range', {}).get('end')}")
        check(not ABS_PATH_RE.search(json.dumps(task, ensure_ascii=False)),
              "建任务响应无服务器绝对路径")
        check("excel_path" not in json.dumps(task), "建任务响应不给 excel_path")

        print("\n=== 3) 重读任务（模拟前端刷新页面）===")
        response = httpx.get(f"{base}/api/tasks/{task_id}", timeout=60)
        reread = response.json()
        print(f"   HTTP {response.status_code} | run_count={reread.get('run_count')} | "
              f"spec_id={reread.get('spec_id')} | schema_version={reread.get('schema_version')}")
        check(response.status_code == 200, "GET 单任务 200", str(response.status_code))
        check(reread.get("spec") == spec, "重读的 Spec 与建任务时**逐字段一致**（不靠前端记内存）")
        check(reread.get("run_count") == 0, "还没跑过 → run_count=0", str(reread.get("run_count")))
        check(reread.get("run_url", "").endswith(f"/api/tasks/{task_id}/run"), "给了 run_url")
        check(reread.get("runs_url", "").endswith(f"/api/tasks/{task_id}/runs"), "给了 runs_url")

        print("\n=== 4) 执行任务（真算数 + 真渲染）===")
        started = time.time()
        response = httpx.post(f"{base}/api/tasks/{task_id}/run", timeout=900)
        run = response.json()
        print(f"   HTTP {response.status_code} | execution_id={run.get('execution_id')} | status={run.get('status')}")
        print(f"   amount={run.get('amount')} amount_full={run.get('amount_full')} "
              f"rows_in_range={run.get('rows_in_range')} rows_excluded={run.get('rows_excluded')} | "
              f"{time.time() - started:.1f}s")
        check(response.status_code == 200, "run HTTP 200", str(response.status_code))
        check(run.get("status") == "success", "执行成功", str(run.get("status")))
        check(abs(float(run.get("amount_full", 0)) - EXPECTED_AMOUNT) < 0.01,
              f"金额 {EXPECTED_AMOUNT}", str(run.get("amount_full")))
        check(run.get("amount") == "316412.16", "展示金额 2 位小数字符串", str(run.get("amount")))
        check(run.get("rows_in_range") == EXPECTED_ROWS_IN_RANGE, f"区间行数 {EXPECTED_ROWS_IN_RANGE}",
              str(run.get("rows_in_range")))
        check(run.get("rows_excluded") == EXPECTED_ROWS_EXCLUDED, f"排除行数 {EXPECTED_ROWS_EXCLUDED}",
              str(run.get("rows_excluded")))
        check(run.get("download_url") == f"/api/download/{run.get('execution_id')}", "给了相对 download_url")
        check(not ABS_PATH_RE.search(json.dumps(run, ensure_ascii=False)), "run 响应无服务器绝对路径")
        check(bool(run.get("code_version")), "记录了代码版本", str(run.get("code_version")))
        check(bool(run.get("verification")), "记录了自校验结果", str(run.get("verification")))
        check(run.get("resolved_range") == {"start": START, "end": END},
              "resolved_range = 真正用于取数的区间", str(run.get("resolved_range")))

        print("\n=== 5) 真下载 + openpyxl 读回 ===")
        response = httpx.get(f"{base}{run['download_url']}", timeout=300)
        data = response.content
        print(f"   HTTP {response.status_code} | {len(data)} 字节 | sha256={hashlib.sha256(data).hexdigest()[:16]}…")
        check(response.status_code == 200, "下载 HTTP 200", str(response.status_code))
        tmp = os.path.join(ROOT, "outputs", "_task_smoke.xlsx")
        with open(tmp, "wb") as handle:
            handle.write(data)
        from openpyxl import load_workbook
        sheet = load_workbook(tmp).active
        cells = {}
        for row in sheet.iter_rows(values_only=True):
            if row and row[0] and row[1] is not None:
                cells[str(row[0])] = row[1]
        print(f"   读回：销售额={cells.get('销售额（元）')} 区间原始行数={cells.get('区间原始行数')} "
              f"被排除行数={cells.get('被排除行数')}")
        check(abs(float(cells.get("销售额（元）", 0)) - EXPECTED_AMOUNT) < 0.01,
              "读回 Excel 的销售额与接口一致", str(cells.get("销售额（元）")))
        check(cells.get("区间原始行数") == EXPECTED_ROWS_IN_RANGE, "读回 Excel 的区间行数一致")
        amount_format = sheet["B4"].number_format
        print(f"   B4 单元格：值={sheet['B4'].value} 数字格式={amount_format!r}")
        check(amount_format == "#,##0.00", "金额单元格保留 2 位小数格式（D17-2/D5）", repr(amount_format))
        os.remove(tmp)

        print("\n=== 6) 任务执行历史 ===")
        response = httpx.get(f"{base}/api/tasks/{task_id}/runs", timeout=60)
        history = response.json()
        runs = history.get("runs") or []
        print(f"   HTTP {response.status_code} | count={history.get('count')} total={history.get('total')} "
              f"limit={history.get('limit')} offset={history.get('offset')}")
        check(history.get("total") == 1, "该任务恰好 1 次执行", str(history.get("total")))
        required_keys = {
            "run_id", "execution_id", "task_id", "status", "created_at", "spec_id", "spec_version",
            "time_range", "resolved_range", "base_date", "data_sha256", "code_version", "verification",
            "seconds", "error", "stage", "output_rel_path", "download_url", "amount", "amount_full",
            "rows_in_range", "rows_valid", "rows_excluded",
        }
        if runs:
            missing = required_keys - set(runs[0])
            check(not missing, "run 记录字段齐全（R005 #6 钉死的清单）", f"缺 {sorted(missing)}" if missing else "")
            check(runs[0].get("task_id") == task_id, "run 记录绑定了 task_id")
            check(runs[0].get("base_date") is None, "绝对区间 → base_date 为 null（相对时间才有）")

        print("\n=== 7) 任务列表（分页 + 统计）===")
        response = httpx.get(f"{base}/api/tasks", params={"limit": 50, "offset": 0}, timeout=60)
        listing = response.json()
        mine = next((item for item in listing.get("tasks", []) if item.get("task_id") == task_id), None)
        print(f"   HTTP {response.status_code} | count={listing.get('count')} total={listing.get('total')}")
        check(mine is not None, "列表里能找到刚建的任务")
        if mine:
            print(f"   该任务：status={mine.get('status')} run_count={mine.get('run_count')} "
                  f"last_run_status={mine.get('last_run_status')} "
                  f"spec_summary.start={mine.get('spec_summary', {}).get('start')}")
            check(mine.get("status") == "has_run", "跑成功后状态 → has_run", str(mine.get("status")))
            check(mine.get("run_count") == 1, "run_count=1", str(mine.get("run_count")))
            check(mine.get("last_run_status") == "success", "last_run_status=success")
            check(mine.get("spec_summary", {}).get("start") == START, "列表带 Spec 摘要（不用拉全量 Spec）")

        print("\n=== 8) 错误路径：未知 task_id ===")
        response = httpx.post(f"{base}/api/tasks/task_不存在/run", timeout=60)
        unknown = response.json()
        print(f"   HTTP {response.status_code} | code={unknown.get('error', {}).get('code')}")
        check(response.status_code == 404, "未知 task → 404", str(response.status_code))
        check(unknown.get("error", {}).get("code") == "task_not_found", "机器可读错误码 task_not_found")

        print("\n=== 9) 错误路径：数据不是冻结快照 → 拒绝执行 + 留失败记录（D10）===")
        foreign_path = os.path.join(sandbox, "_task_smoke_foreign.xlsx")
        from openpyxl import Workbook
        book = Workbook()
        page = book.active
        page.append(["InvoiceNo", "StockCode", "Description", "Quantity", "InvoiceDate", "UnitPrice", "CustomerID", "Country"])
        page.append(["A1", "S1", "冒烟假数据", 3, "2011-11-22 10:00:00", 2.5, 10001, "United Kingdom"])
        book.save(foreign_path)
        with open(foreign_path, "rb") as handle:
            response = httpx.post(f"{base}/api/upload", files={"file": ("foreign.xlsx", handle)}, timeout=300)
        foreign = response.json()
        print(f"   上传非快照文件 → HTTP {response.status_code} | rows={foreign.get('rows')}")
        response = httpx.post(
            f"{base}/api/tasks",
            json={"file_id": foreign["file_id"], "start": START, "end": END, "metrics": ["sales_amount"]},
            timeout=300,
        )
        foreign_task = response.json()
        warnings = foreign_task.get("warnings") or []
        print(f"   建任务 → HTTP {response.status_code} | warnings={len(warnings)} 条")
        check(response.status_code == 201, "非快照文件**能建**任务（建时不拦）", str(response.status_code))
        check(bool(warnings), "建任务时给出警告（建时不拦，执行才拦）")
        runs_before = httpx.get(f"{base}/api/tasks/{foreign_task['task_id']}/runs", timeout=60).json().get("total")
        produced_before = set(os.listdir(ENV["SRA_OUTPUT_DIR"]))
        response = httpx.post(f"{base}/api/tasks/{foreign_task['task_id']}/run", timeout=900)
        rejected = response.json()
        print(f"   执行 → HTTP {response.status_code} | code={rejected.get('error', {}).get('code')}")
        check(response.status_code == 422, "非快照文件执行被拒（422）", str(response.status_code))
        check(rejected.get("error", {}).get("code") == "data_snapshot_mismatch", "错误码 data_snapshot_mismatch")
        runs_after = httpx.get(f"{base}/api/tasks/{foreign_task['task_id']}/runs", timeout=60).json()
        check(runs_after.get("total") == (runs_before or 0) + 1, "被拒也**留了失败记录**（D10：失败要留痕）")
        failed_run = (runs_after.get("runs") or [{}])[0]
        check(failed_run.get("status") == "failed", "记录 status=failed")
        check(failed_run.get("output_rel_path") is None, "失败记录**没有产出文件**（D10）")
        check(failed_run.get("download_url") is None, "失败记录没有 download_url（不可下载）")
        task_after = httpx.get(f"{base}/api/tasks/{foreign_task['task_id']}", timeout=60).json()
        check(task_after.get("status") == "created", "任务状态**仍是 created**（失败不污染任务定义）",
              str(task_after.get("status")))
        check(task_after.get("run_count") == 1, "但 run_count 记到了 1 次（失败也是执行事实）")
        produced_after = set(os.listdir(ENV["SRA_OUTPUT_DIR"]))
        leaked = produced_after - produced_before
        check(not leaked, "被拒的执行**没有落任何新产出文件**（D10：失败不产出）", str(sorted(leaked)))
        os.remove(foreign_path)
        print(f"   沙箱产出目录：{sorted(produced_after)}（只有第 4 步那次成功执行的产出）")

        print("\n=== 10) 统一错误体（新老端点一致）===")
        response = httpx.post(f"{base}/api/execute", json={"file_id": "不存在", "start": START, "end": END}, timeout=60)
        body = response.json()
        print(f"   老端点 /api/execute → HTTP {response.status_code} | code={body.get('error', {}).get('code')}")
        check(response.status_code in (404, 410, 422), "老端点仍返回 4xx", str(response.status_code))
        check(body.get("error", {}).get("code") == "upload_not_found", "老端点也有机器可读 code")
        check("detail" in body, "老端点保留顶层 detail（002C 契约不变，R005 #10）")
    finally:
        proc.terminate()
        try:
            proc.wait(timeout=10)
        except Exception:
            proc.kill()
        print("(service 已关闭)")

    print()
    if failures:
        print(f"❌ 任务层冒烟有 {len(failures)} 项失败：")
        for label in failures:
            print(f"   · {label}")
        return 1
    print("🎉 任务层冒烟全部通过：建任务 → 重读 Spec 一致 → 真执行 316412.16/19950 行 → "
          "真下载可读回 2 位小数 → 记录留痕 → 状态 created→has_run → 非快照数据被拒且不产文件")
    return 0


if __name__ == "__main__":
    sys.exit(main())
