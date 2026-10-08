"""export_docs.py · 报告的三种导出文件（Word / Excel / Markdown）。

════════════════════════════════════════════════════════════════════════
【这个文件负责什么】
════════════════════════════════════════════════════════════════════════
报告的结构（标题 / 分节 / 表格 / 结论建议）由 `report.py` 算出来、**冻结**在会话记录里；
本文件只把它**排成三种真文件**：

    · Word   (.docx) —— 文档形态：标题层级 + 真表格（python-docx）
    · Excel  (.xlsx) —— 表格形态：多工作表 + 数字格式（openpyxl）
    · Markdown(.md)  —— 文本形态：沿用既有的 Markdown 渲染（页面预览的也是那段文本）

**这里一个数字都不算**：所有数值都是从冻结报告里搬过来的。三种格式、页面预览、
以及日志里记着的那份结果，逐位相同 —— 没有第二条计算路径。

【为什么默认是 Word】
用户实测反馈：默认导出 .md，Windows 双击打不开（没有关联程序），体感就是"下载不了"；
而且周报本身是**文档形态**（标题 / 摘要 / 分节 / 表格 / 结论建议），Word 才贴它的样子。
Excel 留给"要拿数字继续算"的场合（多工作表 + 真数字 + 数字格式）；Markdown 仍可选，不再是默认。

【为什么不做 PDF】
本机没有 PDF 库，装一个新的就是增加依赖（项目铁律 7：依赖清单必须精简）；
自己拼 PDF 字节流属于"造轮子"，同样不做。Word 与 Excel 已经覆盖了"能打开、能看、能再算"。

【零新依赖】
`python-docx`（文档输入那条链路在用）与 `openpyxl`（Excel 渲染器在用）都是既有依赖。
"""

from __future__ import annotations

import io
import re
from typing import Any

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Pt
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

# ════════════════════════════════════════════════════════════════════════
# 三种格式的声明（前端只认 label / filename / mime，不自己拼文件类型）
# ════════════════════════════════════════════════════════════════════════
FORMAT_ORDER = ("docx", "xlsx", "md")
DEFAULT_FORMAT = "docx"
FORMATS: dict[str, dict[str, str]] = {
    "docx": {
        "label": "Word 文档",
        "ext": "docx",
        "mime": "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
    },
    "xlsx": {
        "label": "Excel 工作簿",
        "ext": "xlsx",
        "mime": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    },
    "md": {
        "label": "Markdown 文本",
        "ext": "md",
        "mime": "text/markdown;charset=utf-8",
    },
}

# Excel 里的数字格式（写了数字就得给格式，否则用户看到的是 1234.5 这种裸值）
MONEY_FORMAT = "#,##0.00"
SIGNED_MONEY_FORMAT = "+#,##0.00;-#,##0.00"
INT_FORMAT = "#,##0"
SIGNED_INT_FORMAT = "+#,##0;-#,##0"
PERCENT_FORMAT = "+0.00%;-0.00%;0.00%"

# 只认**程序自己生成的展示数字**（千分位 / 两位小数 / 带元或百分号）：
# 日期区间、名次（#1）、说明列都认不出来，就按文本原样写 —— 宁可当文本，也不瞎转。
_PERCENT_CELL = re.compile(r"^([+-]?\d{1,3}(?:,\d{3})*\.\d+)%$")
_MONEY_CELL = re.compile(r"^([+-]?\d{1,3}(?:,\d{3})*(?:\.\d+)?)元?$")


def base_filename(title: str, period_label: str) -> str:
    """文件名的"主干"（不带扩展名）：`销售周报-2011-11-28至2011-12-04`。

    用「至」而不是下划线 —— 这是一份要发出去的文档，不是内部文件名。
    """
    span = str(period_label or "").replace(" ~ ", "至").replace("~", "至").strip()
    return f"{title}-{span}" if span else str(title)


def format_choices(base: str) -> list[dict[str, str]]:
    """发给前端的格式清单（文件名由后端拼，前端一个字都不拼）。"""
    return [
        {
            "format": key,
            "label": FORMATS[key]["label"],
            "filename": f"{base}.{FORMATS[key]['ext']}",
            "mime": FORMATS[key]["mime"],
        }
        for key in FORMAT_ORDER
    ]


def mime_of(fmt: str) -> str:
    return FORMATS.get(fmt, FORMATS[DEFAULT_FORMAT])["mime"]


# ════════════════════════════════════════════════════════════════════════
# 报告里那几段文字：三种格式共用同一套取法
# ════════════════════════════════════════════════════════════════════════
def _section_lines(section: dict[str, Any]) -> str:
    """把「摘要」那种一整段的话拼成一段（`style=paragraph` 的分节）。"""
    return "".join(str(item) for item in section.get("lines") or [] if str(item).strip()).strip()


def _bullet_lines(section: dict[str, Any]) -> list[str]:
    return [str(item).strip() for item in section.get("lines") or [] if str(item).strip()]


def _fact_lines(profile: dict[str, Any]) -> list[str]:
    """文档/工作簿末尾都要写清楚：数字来自哪段数据、金额什么单位（不靠页面上的提示）。"""
    currency = (profile or {}).get("currency") or {}
    lines = []
    if profile.get("first_day") and profile.get("last_day"):
        lines.append(f"数据范围：{profile['first_day']} ~ {profile['last_day']}")
    if currency.get("name"):
        lines.append(
            f"金额单位：{currency['name']}（{currency.get('code', '')}）；"
            "币种由项目唯一声明给出，数值未做任何换算。"
        )
    return lines


def _source_note(report: dict[str, Any]) -> str:
    return (
        f"比较口径：{report.get('comparison_label', '—')}；"
        "报告里的数字全部由程序确定性计算得出（两区间比较 / 汇总 / 趋势 / 国家分布 / "
        "商品排行五个既有计算能力），数字逐项核对过。"
    )


# ════════════════════════════════════════════════════════════════════════
# Word（.docx）
# ════════════════════════════════════════════════════════════════════════
# 中文字体要单独指一次「东亚字体」：python-docx 的默认模板是 Calibri，
# 只设 font.name 的话中文会走系统回退，标题与表格的观感会不一致。
_CN_FONT = "微软雅黑"
_CN_STYLES = ("Normal", "Title", "Heading 1", "Heading 2", "Heading 3", "List Bullet")


def _apply_cn_font(document: Document) -> None:
    for name in _CN_STYLES:
        try:
            style = document.styles[name]
        except KeyError:                     # 模板里没有这个样式就跳过（不硬造）
            continue
        style.font.name = _CN_FONT
        rpr = style.element.get_or_add_rPr()
        rpr.get_or_add_rFonts().set(qn("w:eastAsia"), _CN_FONT)


def _add_table(document: Document, table: dict[str, Any]) -> None:
    """真表格（不是把表格拍成文本）：表头加粗、数字右对齐。"""
    headers = [str(item) for item in table.get("headers") or []]
    rows = table.get("rows") or []
    if not headers:
        return
    grid = document.add_table(rows=1, cols=len(headers))
    grid.style = "Table Grid"
    for index, head in enumerate(headers):
        cell = grid.rows[0].cells[index]
        cell.text = ""
        run = cell.paragraphs[0].add_run(head)
        run.bold = True
        run.font.size = Pt(9)
    for row in rows:
        cells = grid.add_row().cells
        for index, value in enumerate(row[: len(headers)]):
            text = str(value)
            cells[index].text = ""
            run = cells[index].paragraphs[0].add_run(text)
            run.font.size = Pt(9)
            if _MONEY_CELL.match(text.strip()) or _PERCENT_CELL.match(text.strip()):
                cells[index].paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.RIGHT
    document.add_paragraph()


def render_docx(
    report: dict[str, Any],
    *,
    why_text: str,
    actions_text: str,
    profile: dict[str, Any],
) -> bytes:
    """报告 → Word 文档（标题层级 + 真表格 + 结论建议 + 口径脚注）。"""
    document = Document()
    _apply_cn_font(document)
    document.add_heading(f"{report.get('title', '销售报告')} · {report.get('period_label', '')}", level=0)
    document.add_paragraph(_source_note(report))

    for section in report.get("sections") or []:
        document.add_heading(str(section.get("title") or ""), level=1)
        for block in section.get("blocks") or []:
            document.add_heading(str(block.get("title") or ""), level=2)
            _add_table(document, block.get("table") or {})
        if section.get("table"):
            _add_table(document, section["table"])
        if section.get("style") == "paragraph":
            paragraph = _section_lines(section)
            if paragraph:
                document.add_paragraph(paragraph)
        else:
            for item in _bullet_lines(section):
                document.add_paragraph(item, style="List Bullet")

    document.add_heading("结论与建议", level=1)
    for title, text in (("【为什么】", why_text), ("【建议行动】", actions_text)):
        document.add_heading(title, level=2)
        document.add_paragraph(str(text or "").strip())

    for line in _fact_lines(profile):
        note = document.add_paragraph(line)
        note.runs[0].font.size = Pt(8)
        note.runs[0].italic = True

    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


# ════════════════════════════════════════════════════════════════════════
# Excel（.xlsx）
# ════════════════════════════════════════════════════════════════════════
_HEAD_FILL = PatternFill("solid", fgColor="DCE6F1")
_TITLE_FONT = Font(bold=True, size=14)
_HEAD_FONT = Font(bold=True, size=10)

# 工作表 ← 报告分节：`(工作表名, 分节 key, 块标题前缀)`
# 前两个直接搬分节里的那张 table；后两个从「结构」那一节的 blocks 里按标题各取一块。
SHEET_SECTIONS = (
    ("核心指标", "metrics", ""),
    ("趋势", "trend", ""),
    ("结构-国家", "structure", "国家分布"),
    ("结构-商品", "structure", "商品"),
)


def _cell(text: Any) -> tuple[Any, str | None]:
    """展示串 → (单元格值, 数字格式)。

    **只做无损转换**：`1,234.56元` → 数字 1234.56 + `#,##0.00`（显示出来一模一样），
    认不出来的（日期区间 / `#1` / 说明列）一律按文本原样写 —— 不猜、不重算。
    这样工作表里的数字能直接拿去算，而它和页面上的还是同一个数。
    """
    raw = str(text if text is not None else "").strip()
    if not raw:
        return "", None
    percent = _PERCENT_CELL.match(raw)
    if percent:
        return round(float(percent.group(1).replace(",", "")) / 100, 8), PERCENT_FORMAT
    money = _MONEY_CELL.match(raw)
    if money:
        number = float(money.group(1).replace(",", ""))
        signed = raw[0] in "+-"
        has_decimals = "." in money.group(1)
        if has_decimals:
            return number, SIGNED_MONEY_FORMAT if signed else MONEY_FORMAT
        return int(number), SIGNED_INT_FORMAT if signed else INT_FORMAT
    return raw, None


def _write_table(sheet, table: dict[str, Any], start_row: int, *, title: str = "") -> int:
    """把一张表写到工作表上（表头加粗带底色、数字右对齐、列宽按内容给）。"""
    headers = [str(item) for item in table.get("headers") or []]
    rows = table.get("rows") or []
    if not headers:
        return start_row
    row = start_row
    if title:
        sheet.cell(row=row, column=1, value=title).font = Font(bold=True, size=11)
        row += 1
    for index, head in enumerate(headers, start=1):
        cell = sheet.cell(row=row, column=index, value=head)
        cell.font = _HEAD_FONT
        cell.fill = _HEAD_FILL
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    row += 1
    for data in rows:
        for index, value in enumerate(data[: len(headers)], start=1):
            number, number_format = _cell(value)
            cell = sheet.cell(row=row, column=index, value=number)
            if number_format:
                cell.number_format = number_format
                cell.alignment = Alignment(horizontal="right")
            else:
                cell.alignment = Alignment(horizontal="left", wrap_text=False)
        row += 1
    return row + 1


def _autosize(sheet, columns: int) -> None:
    """列宽按内容长度给（中文按两个字符宽算），上限 60 —— 不用打开就手动拖列。"""
    for index in range(1, columns + 1):
        width = 8
        for cell in sheet[get_column_letter(index)]:
            text = str(cell.value or "")
            width = max(width, sum(2 if ord(char) > 0x2E80 else 1 for char in text) + 2)
        sheet.column_dimensions[get_column_letter(index)].width = min(width, 60)


def render_xlsx(
    report: dict[str, Any],
    *,
    why_text: str,
    actions_text: str,
    profile: dict[str, Any],
) -> bytes:
    """报告 → Excel 工作簿：摘要与建议 / 核心指标 / 趋势 / 结构（国家）/ 结构（商品）。"""
    workbook = Workbook()

    # ── ① 摘要与建议（文档性的内容放在第一张表）────────────────────────
    summary_sheet = workbook.active
    summary_sheet.title = "摘要与建议"
    summary = next((item for item in report.get("sections") or []
                    if item.get("key") == "summary"), {})
    row = 1
    summary_sheet.cell(row=row, column=1,
                       value=f"{report.get('title', '销售报告')} · {report.get('period_label', '')}").font = _TITLE_FONT
    row += 2
    for line in [_source_note(report), *_fact_lines(profile)]:
        summary_sheet.cell(row=row, column=1, value=line)
        row += 1
    row += 1
    summary_sheet.cell(row=row, column=1, value="摘要").font = _HEAD_FONT
    row += 1
    text = _section_lines(summary) or "（本次没有摘要）"
    summary_sheet.cell(row=row, column=1, value=text).alignment = Alignment(wrap_text=True, vertical="top")
    summary_sheet.merge_cells(start_row=row, start_column=1, end_row=row, end_column=4)
    row += 2
    for title, body in (("【为什么】", why_text), ("【建议行动】", actions_text)):
        summary_sheet.cell(row=row, column=1, value=title).font = _HEAD_FONT
        row += 1
        summary_sheet.cell(row=row, column=1, value=str(body or "").strip()).alignment = Alignment(
            wrap_text=True, vertical="top")
        summary_sheet.merge_cells(start_row=row, start_column=1, end_row=row, end_column=4)
        row += 2
    summary_sheet.column_dimensions["A"].width = 96

    # ── ② 后面几张表：一张工作表放一张表（要拿去继续算的数字都在这几张里）──
    # 工作表 ↔ 分节的对应关系**写死**（报告结构改了这里要跟着改）：不靠标题猜，
    # 猜错了会把两张表写进同一张工作表，而且不报错。
    sections_by_key = {section.get("key"): section for section in report.get("sections") or []}
    for sheet_name, section_key, block_prefix in SHEET_SECTIONS:
        sheet = workbook.create_sheet(sheet_name)
        section = sections_by_key.get(section_key) or {}
        line = 1
        if block_prefix:
            # 结构那一节分成两个 block（国家 / 商品），按块标题各写一张工作表
            for block in section.get("blocks") or []:
                if str(block.get("title") or "").startswith(block_prefix):
                    line = _write_table(sheet, block.get("table") or {}, line,
                                        title=str(block.get("title") or ""))
        else:
            line = _write_table(sheet, section.get("table") or {}, line,
                                title=str(section.get("title") or ""))
        if sheet.max_row == 1 and sheet.max_column == 1:
            sheet.cell(row=1, column=1, value="本次报告没有这张表。")   # 空表也说一句，不留白板
        _autosize(sheet, max(1, sheet.max_column))

    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


# ════════════════════════════════════════════════════════════════════════
# 统一入口
# ════════════════════════════════════════════════════════════════════════
def render_report(
    report: dict[str, Any],
    fmt: str,
    *,
    why_text: str,
    actions_text: str,
    profile: dict[str, Any],
    markdown: str = "",
) -> tuple[bytes, str, str]:
    """按格式渲染 → `(文件内容, 文件名, 内容类型)`。

    Markdown 那一路直接用调用方给的文本（就是页面上预览的那段，一个字都不重排）；
    Word / Excel 从冻结报告渲染 —— 三种格式的数字都来自同一份结果。
    """
    key = fmt if fmt in FORMATS else DEFAULT_FORMAT
    filename = f"{base_filename(report.get('title', '销售报告'), report.get('period_label', ''))}." \
               f"{FORMATS[key]['ext']}"
    if key == "docx":
        return render_docx(report, why_text=why_text, actions_text=actions_text, profile=profile), \
            filename, FORMATS[key]["mime"]
    if key == "xlsx":
        return render_xlsx(report, why_text=why_text, actions_text=actions_text, profile=profile), \
            filename, FORMATS[key]["mime"]
    return (markdown or "").encode("utf-8"), filename, FORMATS[key]["mime"]


__all__ = [
    "DEFAULT_FORMAT",
    "FORMATS",
    "FORMAT_ORDER",
    "base_filename",
    "format_choices",
    "mime_of",
    "render_report",
]
