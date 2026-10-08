"""tests/test_fr002b_desktop.py · FR-002B：导出直达桌面 + 「期间」列日期序列号缺陷。

════════════════════════════════════════════════════════════════════════
【用户的原话与真因】
════════════════════════════════════════════════════════════════════════
> 「在提问如『帮我生成一星期以内的周报』的时候，虽然能够生成 word / excel / markdown
>   等数据，但是**还是那个问题，无法打开**。**你帮我弄到桌面**。」

实测出来的真因是两条，都不是"文件坏了"：

  ① **文件掉在了别处**：用户平时用的 Edge 配置里下载目录被设成了 `D:\\`，
     于是下载下来的文件全在 D 盘根目录，用户在桌面上当然找不到 —— 体感就是"打不开"。
     → 所以要让服务端**直接把文件写到桌面**，并且能一键在资源管理器里定位它。
  ② **「期间」列导出成了日期序列号**：`销售数据-销售表-*.xlsx` 的第一列显示 `40875`
     而不是 `2011-11-28`。因为那一列声明的是文本格式（`@`），而单元格里塞的是
     日期值 —— Excel 拿到"日期值 + 文本格式"这对组合，就把内部序列号原样显示出来。

════════════════════════════════════════════════════════════════════════
【本文件怎么验（真文件、真读回，不看"Excel 打开像是正常的"）】
════════════════════════════════════════════════════════════════════════
· 桌面落到**临时目录**（`SRA_DESKTOP_DIR`），绝不碰真实桌面；
· 导出物用 `openpyxl` **读回来**，断言 `cell.value` 的类型与 `cell.number_format`
  （评审要求：不能只检查"看起来正常"）；
· 「在文件夹中打开」把 `subprocess.Popen` 换掉，断言**真的只调了 explorer.exe**，
  且路径来自服务端自己的记录（不是请求里带来的）。
"""

from __future__ import annotations

import datetime as _dt
import inspect
import io
import os
import pathlib
import sys

import pytest

PROJECT_ROOT = pathlib.Path(__file__).resolve().parents[1]
if str(PROJECT_ROOT) not in sys.path:
    sys.path.insert(0, str(PROJECT_ROOT))

from fastapi.testclient import TestClient  # noqa: E402
from openpyxl import load_workbook  # noqa: E402
from openpyxl.styles.numbers import is_date_format  # noqa: E402

from app import desktop  # noqa: E402
from app.ai import llm  # noqa: E402
from app.api import app  # noqa: E402

client = TestClient(app)

WEEKLY_QUESTION = "帮我根据本星期的销售数据做一份销售周报"
# 一个 5 天的区间：够看出「期间」是逐日的日期，又不至于让导出文件很大
SALES_QUERY = {"table": "sales", "format": "xlsx", "dimension": "day",
               "start": "2011-11-01", "end": "2011-11-05"}


@pytest.fixture(autouse=True)
def isolated_dirs(tmp_path, monkeypatch):
    """所有出口都指到临时目录 —— **真实桌面一个字节都不写**。"""
    monkeypatch.setenv("SRA_STATE_DIR", str(tmp_path / "state"))
    monkeypatch.setenv("SRA_DOC_DIR", str(tmp_path / "documents"))
    monkeypatch.setenv("SRA_UPLOAD_DIR", str(tmp_path / "uploads"))
    monkeypatch.setenv("SRA_OUTPUT_DIR", str(tmp_path / "outputs"))
    monkeypatch.setenv(desktop.DESKTOP_DIR_ENV, str(_desktop_dir(tmp_path)))
    desktop.reset_last_saved()
    return tmp_path


@pytest.fixture(autouse=True)
def no_llm(monkeypatch):
    """关掉模型通道：报告结论段走代码降级文案，导出照样成立。"""
    monkeypatch.setattr(llm, "available", lambda: False)
    return None


def _desktop_dir(tmp_path: pathlib.Path) -> pathlib.Path:
    directory = tmp_path / "Desktop"
    directory.mkdir(exist_ok=True)
    return directory


def _save(payload: dict):
    return client.post("/api/exports/desktop", json=payload)


def _code_only(module) -> str:
    """模块的**可执行代码**（注释与文档字符串都去掉）—— 用于"这条写法不许出现"的源码检查。"""
    import ast

    tree = ast.parse(inspect.getsource(module))
    for node in ast.walk(tree):
        if not isinstance(node, (ast.Module, ast.FunctionDef, ast.AsyncFunctionDef, ast.ClassDef)):
            continue
        body = node.body
        if body and isinstance(body[0], ast.Expr) and isinstance(body[0].value, ast.Constant) \
                and isinstance(body[0].value.value, str):
            node.body = body[1:]                     # 丢掉文档字符串
    return ast.unparse(tree)


def _saved_path(body: dict) -> pathlib.Path:
    path = pathlib.Path(body["desktop_dir"]) / body["file_name"]
    assert path.is_file(), f"响应说存好了，但文件不在：{path}"
    return path


# ════════════════════════════════════════════════════════════════════════
# 1. 桌面目录：走系统机制解析（不是拼字符串）
# ════════════════════════════════════════════════════════════════════════
def test_C4_桌面目录禁用home与硬编码写法():
    """`Path.home() / "Desktop"` 与硬编码某人的桌面路径，都是本项目**明确不用**的做法。

    中文 Windows、OneDrive 重定向、组策略改桌面位置 —— 只要不是"默认英文、没重定向"，
    这两种写法就都是错的。这条直接钉可执行代码（注释与文档字符串里可以提到它们——
    那里正是写"为什么不用"的地方）。
    """
    code = _code_only(desktop)
    assert ".home()" not in code, "代码里用 home() 拼桌面路径（OneDrive/中文 Windows 会翻车）"
    assert "Users\\\\" not in code and "C:/Users" not in code, "硬编码了某个用户的桌面路径"
    # 三个系统机制都在，且顺序即优先级（权威的排前面）
    names = [probe.__name__ for probe in desktop._SYSTEM_PROBES]
    assert names == ["_known_folder_desktop", "_registry_desktop", "_shfolder_desktop"]
    assert "SHGetKnownFolderPath" in inspect.getsource(desktop._known_folder_desktop)
    assert "winreg" in inspect.getsource(desktop._registry_desktop)
    assert "SHGetFolderPathW" in inspect.getsource(desktop._shfolder_desktop)


def test_C4_主机制在本机真的能给出答案():
    """主机制（外壳 API）必须**真的返回路径**，不是静默 None。

    实测踩过：`SHGetKnownFolderPath` 导出在 **shell32.dll**，写成 `ole32.dll` 时
    ctypes 直接抛 `function not found`，被 `except` 吞掉 → 主机制永远失效、
    悄悄退到次选机制。那样在 OneDrive 重定向下就可能拿不到真正的桌面，所以钉死这条。
    """
    if sys.platform != "win32":
        pytest.skip("外壳 API 只在 Windows 上有")
    assert desktop._known_folder_desktop() is not None, "主机制没给出桌面目录（机制链等于降级了）"


def test_C4_系统机制给出的目录原样采用(tmp_path, monkeypatch):
    """探测返回什么就用什么，不做任何"猜一个 Desktop 子目录"的加工。

    这里模拟的正是 OneDrive 重定向：桌面其实在 `…\\OneDrive\\Desktop`，
    只有问系统才知道；拼接写法会指向一个**不存在或已废弃**的旧桌面。
    """
    monkeypatch.delenv(desktop.DESKTOP_DIR_ENV, raising=False)
    onedrive = tmp_path / "OneDrive" / "Desktop"
    onedrive.mkdir(parents=True)
    got = desktop.resolve_desktop_dir(probes=(lambda: onedrive,))
    assert got == onedrive


def test_C4_探测链按顺序回退(tmp_path, monkeypatch):
    """第一个机制失败（抛异常）、第二个不给答案（None）→ 用第三个；顺序即优先级。"""
    monkeypatch.delenv(desktop.DESKTOP_DIR_ENV, raising=False)
    real = tmp_path / "RealDesktop"
    real.mkdir()

    def boom():
        raise OSError("外壳 API 不可用")

    got = desktop.resolve_desktop_dir(probes=(boom, lambda: None, lambda: real))
    assert got == real


def test_C4_返回的不是目录就跳过(tmp_path, monkeypatch):
    """机制回了一个"路径"但它不是目录（例如被挪走的旧桌面）→ 不能当真。"""
    monkeypatch.delenv(desktop.DESKTOP_DIR_ENV, raising=False)
    file_like = tmp_path / "not-a-dir.txt"
    file_like.write_text("x", encoding="utf-8")
    real = tmp_path / "RealDesktop"
    real.mkdir()
    assert desktop.resolve_desktop_dir(probes=(lambda: file_like, lambda: real)) == real


def test_C9_拿不到桌面时明确报错且不写到别处(monkeypatch):
    """三个机制都没答案 → 报错。**不许**静默退回当前目录/家目录/输出目录。"""
    monkeypatch.delenv(desktop.DESKTOP_DIR_ENV, raising=False)
    monkeypatch.setattr(desktop, "_SYSTEM_PROBES", ())

    with pytest.raises(desktop.DesktopError) as info:
        desktop.resolve_desktop_dir()
    assert info.value.code == "desktop_unavailable"
    assert "没有" in info.value.message and "别的地方" in info.value.message

    # 落盘这条路同样报错，而且**什么都没记下**（没有"悄悄存到别处"这回事）
    with pytest.raises(desktop.DesktopError) as info2:
        desktop.save_to_desktop("x.xlsx", b"123")
    assert info2.value.code == "desktop_unavailable"
    assert desktop.last_saved() is None


def test_C9_环境变量指向非目录时明确报错(tmp_path, monkeypatch):
    """显式覆盖（SRA_DESKTOP_DIR）也要过 is_dir 验证 —— 覆盖不是"兜底"。"""
    target = tmp_path / "a-file.txt"
    target.write_text("x", encoding="utf-8")
    monkeypatch.setenv(desktop.DESKTOP_DIR_ENV, str(target))
    with pytest.raises(desktop.DesktopError) as info:
        desktop.resolve_desktop_dir()
    assert info.value.code == "desktop_not_a_directory"


def test_接口层拿不到桌面时返回503而不是500(tmp_path, monkeypatch):
    """服务端环境问题（拿不到桌面）与用户输入问题是两码事：503 + 可读中文说明。"""
    monkeypatch.delenv(desktop.DESKTOP_DIR_ENV, raising=False)
    monkeypatch.setattr(desktop, "_SYSTEM_PROBES", ())
    response = _save({"source": "table", **SALES_QUERY})
    assert response.status_code == 503, response.text
    assert response.json()["error"]["code"] == "desktop_unavailable"


# ════════════════════════════════════════════════════════════════════════
# 2. 文件名：清洗注入字符（C6）
# ════════════════════════════════════════════════════════════════════════
def test_C6_文件名清洗禁止注入字符与目录穿越():
    # 目录穿越：**先**只取最后一段（`/` 与 `\` 都算分隔符），路径穿不出去
    assert desktop.sanitize_filename(r"..\..\Windows\System32\evil.xlsx") == "evil.xlsx"
    assert desktop.sanitize_filename("/etc/passwd") == "passwd"
    assert desktop.sanitize_filename("x/y\\z.xlsx") == "z.xlsx"
    assert desktop.sanitize_filename("..") == "导出文件"
    # 禁用字符（不含分隔符）一律换成下划线
    got = desktop.sanitize_filename('a<b>c:d"e|f?g*h.xlsx')
    assert got == "a_b_c_d_e_f_g_h.xlsx"
    assert not set(got) & set('<>:"/\\|?*')
    # 控制字符（含换行）不留
    assert desktop.sanitize_filename("a\nb\tc.xlsx") == "a_b_c.xlsx"
    # 首尾的点和空格去掉（Windows 不许文件名以它们结尾）
    assert desktop.sanitize_filename("  ..x.xlsx..  ") == "x.xlsx"
    # 全是禁用字符 → 落到兜底名
    assert desktop.sanitize_filename("///") == "导出文件"
    assert desktop.sanitize_filename("") == "导出文件"


def test_C6_保留设备名与超长名字():
    # CON / NUL 这类名字在 Windows 上是设备，直接当文件名会失败或被系统吃掉
    assert desktop.sanitize_filename("CON.xlsx") == "_CON.xlsx"
    assert desktop.sanitize_filename("nul.csv") == "_nul.csv"
    long_name = "销售" * 200 + ".xlsx"
    got = desktop.sanitize_filename(long_name)
    assert len(got) <= desktop._MAX_NAME_CHARS and got.endswith(".xlsx")


def test_清洗后落盘的文件名就是响应里的名字(tmp_path):
    """清洗不是只在函数里跑一遍 —— 真正落地的名字以清洗结果为准。"""
    record = desktop.save_to_desktop("报表:2026?28*.xlsx", b"x")
    assert record["file_name"] == "报表_2026_28_.xlsx"
    assert (tmp_path / "Desktop" / record["file_name"]).is_file()


# ════════════════════════════════════════════════════════════════════════
# 3. 不覆盖同名（C5）+ 原子写与无残留（C7）
# ════════════════════════════════════════════════════════════════════════
def test_C5_同名不覆盖_依次排到_1_2(tmp_path):
    desktop_dir = tmp_path / "Desktop"
    first = desktop.save_to_desktop("销售周报-2026-09-28.xlsx", b"one")
    second = desktop.save_to_desktop("销售周报-2026-09-28.xlsx", b"two")
    third = desktop.save_to_desktop("销售周报-2026-09-28.xlsx", b"three")

    assert first["file_name"] == "销售周报-2026-09-28.xlsx"
    assert second["file_name"] == "销售周报-2026-09-28 (1).xlsx"
    assert third["file_name"] == "销售周报-2026-09-28 (2).xlsx"
    # 老文件**内容没被覆盖**（不是"改了名字但还是被写花了"）
    assert (desktop_dir / first["file_name"]).read_bytes() == b"one"
    assert (desktop_dir / second["file_name"]).read_bytes() == b"two"
    assert (desktop_dir / third["file_name"]).read_bytes() == b"three"


def test_C7_落盘后没有临时文件残留(tmp_path):
    desktop_dir = tmp_path / "Desktop"
    desktop.save_to_desktop("x.xlsx", b"content")
    names = sorted(item.name for item in desktop_dir.iterdir())
    assert names == ["x.xlsx"], f"桌面目录里有多余的东西（临时文件/占位文件）：{names}"
    assert not list(desktop_dir.glob("*.part"))


def test_C7_写入失败时清理临时文件且不留占位文件(tmp_path, monkeypatch):
    """改名那一步失败 = 这次导出失败：临时文件与抢到的空占位文件都必须清掉。"""
    desktop_dir = tmp_path / "Desktop"

    def boom(*_args, **_kwargs):
        raise OSError(13, "Permission denied")

    monkeypatch.setattr(desktop.os, "replace", boom)
    with pytest.raises(desktop.DesktopError) as info:
        desktop.save_to_desktop("x.xlsx", b"content")
    assert info.value.code == "desktop_write_failed"

    leftovers = sorted(item.name for item in desktop_dir.iterdir())
    assert leftovers == [], f"失败后桌面目录里留了东西：{leftovers}"
    assert desktop.last_saved() is None, "失败了却记下了'已存到桌面'"


def test_写临时文件失败也清理(tmp_path, monkeypatch):
    desktop_dir = tmp_path / "Desktop"

    def boom(*_args, **_kwargs):
        raise OSError(28, "No space left on device")

    monkeypatch.setattr(desktop.tempfile, "mkstemp", boom)
    with pytest.raises(desktop.DesktopError) as info:
        desktop.save_to_desktop("x.xlsx", b"content")
    assert info.value.code == "desktop_write_failed"
    assert sorted(item.name for item in desktop_dir.iterdir()) == []


# ════════════════════════════════════════════════════════════════════════
# 4. 两条出口：真落盘（C1 / C2）+ 老下载路径仍在（C8）
# ════════════════════════════════════════════════════════════════════════
def test_C1_销售表xlsx真的落到桌面(tmp_path):
    response = _save({"source": "table", **SALES_QUERY})
    assert response.status_code == 201, response.text
    body = response.json()

    path = _saved_path(body)
    assert path.suffix == ".xlsx"
    assert body["bytes"] == path.stat().st_size > 0
    assert body["desktop_dir"] == str(tmp_path / "Desktop")
    assert "同机" in body["deployment"], "响应没写清'服务端与桌面同机'这个前提"
    # 真文件、真能读回（不是零字节占位）
    workbook = load_workbook(path)
    sheet = workbook[workbook.sheetnames[0]]
    assert sheet.max_row >= 2 and sheet.cell(row=1, column=1).value == "期间"


def test_C2_报告docx真的落到桌面(tmp_path):
    chat = client.post("/api/chat", json={"question": WEEKLY_QUESTION, "use_llm": False})
    assert chat.status_code == 200, chat.text
    conversation_id = chat.json()["conversation_id"]

    response = _save({"source": "report", "conversation_id": conversation_id, "format": "docx"})
    assert response.status_code == 201, response.text
    body = response.json()
    path = _saved_path(body)

    assert path.suffix == ".docx" and body["bytes"] > 2000
    from docx import Document

    document = Document(str(path))
    assert document.paragraphs, "落盘的 docx 读回来是空的"
    assert document.tables, "报告表格没写进 docx"


def test_C2_报告xlsx与markdown也能落桌面(tmp_path):
    chat = client.post("/api/chat", json={"question": WEEKLY_QUESTION, "use_llm": False})
    conversation_id = chat.json()["conversation_id"]
    for fmt, suffix in (("xlsx", ".xlsx"), ("md", ".md")):
        body = _save({"source": "report", "conversation_id": conversation_id, "format": fmt}).json()
        path = _saved_path(body)
        assert path.suffix == suffix and path.stat().st_size > 0


def test_落到桌面的内容与下载到的完全一致(tmp_path):
    """同源铁律：两条出口调的是**同一个渲染函数**，不是各渲一遍。

    为什么比单元格而不是比字节：openpyxl 会把"生成时间"写进工作簿属性，
    两次渲染的 zip 字节天然不同（时间戳不一样）—— 比字节会变成比时钟。
    逐格比 `值 + 数字格式`才是"内容相同"的真证据。
    """
    download = client.get("/api/tables/sales/export",
                          params={**SALES_QUERY, "format": "xlsx", "start": "2011-11-01",
                                  "end": "2011-11-05"})
    assert download.status_code == 200
    body = _save({"source": "table", **SALES_QUERY}).json()

    downloaded = load_workbook(io.BytesIO(download.content))
    from_desktop = load_workbook(_saved_path(body))
    assert downloaded.sheetnames == from_desktop.sheetnames

    def cells(workbook):
        sheet = workbook[workbook.sheetnames[0]]
        return [[(cell.value, cell.number_format) for cell in row] for row in sheet.iter_rows()]

    assert cells(downloaded) == cells(from_desktop)


def test_C8_原浏览器下载路径仍然可用():
    """④ 老方案不许因为新方案被删掉：两条老导出端点原地保留。"""
    table_export = client.get("/api/tables/sales/export",
                              params={"format": "xlsx", "dimension": "day",
                                      "start": "2011-11-01", "end": "2011-11-05"})
    assert table_export.status_code == 200
    assert table_export.headers["content-disposition"].startswith("attachment;")
    assert load_workbook(io.BytesIO(table_export.content)).sheetnames

    chat = client.post("/api/chat", json={"question": WEEKLY_QUESTION, "use_llm": False})
    conversation_id = chat.json()["conversation_id"]
    report_export = client.get(f"/api/conversations/{conversation_id}/report/export",
                               params={"format": "docx"})
    assert report_export.status_code == 200 and len(report_export.content) > 2000


def test_C8_导出CSV也照旧():
    response = client.get("/api/tables/sales/export",
                          params={"format": "csv", "dimension": "day",
                                  "start": "2011-11-01", "end": "2011-11-05"})
    assert response.status_code == 200
    text = response.content.decode("utf-8-sig")
    assert text.splitlines()[0].startswith("期间,")


# ════════════════════════════════════════════════════════════════════════
# 5. 请求形状：路径不许由客户端说了算（F1）
# ════════════════════════════════════════════════════════════════════════
def test_F1_请求体带路径字段直接被拒():
    """服务端不接受客户端指定的落盘位置 —— 多一个字段就 422（不是"忽略"）。"""
    for extra in ({"path": "C:\\Windows\\System32\\evil.xlsx"},
                  {"dir": "D:\\"}, {"save_path": "x"}, {"file_path": "x"}):
        response = _save({"source": "table", **SALES_QUERY, **extra})
        assert response.status_code == 422, f"{extra} 居然被接受了"
        assert response.json()["error"]["code"] == "validation_error"


def test_F1_reveal请求体也不接受路径():
    response = client.post("/api/exports/reveal", json={"path": "C:\\Windows\\notepad.exe"})
    assert response.status_code == 422
    assert response.json()["error"]["code"] == "validation_error"


def test_来源不匹配或缺参数时明确报错():
    assert _save({"source": "table", "format": "xlsx"}).json()["error"]["code"] == \
        "export_table_missing"
    assert _save({"source": "report", "format": "docx"}).json()["error"]["code"] == \
        "export_conversation_missing"
    # 两个来源的参数混着传：明说"别混"，不是悄悄忽略
    mixed = _save({"source": "report", "conversation_id": "x", "table": "sales"})
    assert mixed.status_code == 400
    assert mixed.json()["error"]["code"] == "export_source_mismatch"
    assert _save({"source": "table", **SALES_QUERY, "conversation_id": "x"}).status_code == 400


def test_不存在的问答与不支持的格式():
    assert _save({"source": "report", "conversation_id": "nope", "format": "docx"}).status_code == 404
    assert _save({"source": "table", **{**SALES_QUERY, "format": "pdf"}}).status_code == 400


# ════════════════════════════════════════════════════════════════════════
# 6. 「在文件夹中打开」：只认刚刚生成的那个文件
# ════════════════════════════════════════════════════════════════════════
def test_还没存过文件时reveal明确报错():
    response = client.post("/api/exports/reveal", json={})
    assert response.status_code == 404
    assert response.json()["error"]["code"] == "nothing_to_reveal"


def test_reveal只调explorer且路径来自服务端记录(monkeypatch):
    calls: list[list[str]] = []
    monkeypatch.setattr(desktop.subprocess, "Popen", lambda argv: calls.append(argv))

    body = _save({"source": "table", **SALES_QUERY}).json()
    response = client.post("/api/exports/reveal", json={"export_id": body["export_id"]})
    assert response.status_code == 200, response.text
    assert response.json()["revealed"] is True

    assert len(calls) == 1
    argv = calls[0]
    assert argv[0] == "explorer.exe", argv
    # 路径是服务端自己落盘的那一条，且用 /select,"…" 的形式定位到具体文件
    expected = str(pathlib.Path(body["desktop_dir"]) / body["file_name"])
    assert argv[1] == f'/select,"{expected}"', argv


def test_只能定位刚刚生成的那个文件_旧编号被拒(monkeypatch):
    monkeypatch.setattr(desktop.subprocess, "Popen", lambda argv: None)
    first = _save({"source": "table", **SALES_QUERY}).json()
    second = _save({"source": "table", **SALES_QUERY}).json()
    assert first["export_id"] != second["export_id"]

    stale = client.post("/api/exports/reveal", json={"export_id": first["export_id"]})
    assert stale.status_code == 409
    assert stale.json()["error"]["code"] == "export_superseded"
    # 省略编号 = 定位最近那一个
    assert client.post("/api/exports/reveal", json={}).status_code == 200


def test_文件被删掉后reveal如实说找不到(monkeypatch):
    monkeypatch.setattr(desktop.subprocess, "Popen", lambda argv: None)
    body = _save({"source": "table", **SALES_QUERY}).json()
    pathlib.Path(body["desktop_dir"], body["file_name"]).unlink()
    response = client.post("/api/exports/reveal", json={})
    assert response.status_code == 404
    assert response.json()["error"]["code"] == "file_missing"


def test_非Windows平台明确说不支持(monkeypatch):
    """别的平台**没有实现**，就如实返回 501 —— 不假装能用、不静默什么都不做。"""
    monkeypatch.setattr(desktop.sys, "platform", "linux")
    assert desktop.reveal_supported() is False
    body = _save({"source": "table", **SALES_QUERY}).json()
    response = client.post("/api/exports/reveal", json={"export_id": body["export_id"]})
    assert response.status_code == 501
    assert response.json()["error"]["code"] == "reveal_unsupported"


# ════════════════════════════════════════════════════════════════════════
# 7. 「期间」列：真日期（D1 / D2 / D3 / D4 / D5）
# ════════════════════════════════════════════════════════════════════════
def _first_column_cells(path: pathlib.Path, header: str):
    workbook = load_workbook(path)
    sheet = workbook[workbook.sheetnames[0]]
    heads = [cell.value for cell in sheet[1]]
    index = heads.index(header) + 1
    return sheet, [sheet.cell(row=row, column=index)
                   for row in range(2, sheet.max_row + 1)]


@pytest.mark.parametrize("dimension", ["day", "week"])
def test_D1_D2_D3_销售表期间列导出成真日期(dimension, tmp_path):
    """② 的回归测试（评审原话）：原始值 → 导出 → openpyxl 读回 → 是真 date/datetime
    → number_format 是合法日期格式。**不能只检查"Excel 看起来正常"**。"""
    from app.datasets import queries

    body = _save({"source": "table", **{**SALES_QUERY, "dimension": dimension}}).json()
    path = _saved_path(body)
    axis_label = queries.SALES_AXIS_LABELS[dimension]
    _sheet, cells = _first_column_cells(path, axis_label)
    assert cells, "导出的销售表一行数据都没有，这条断言等于没验"

    # 原始值（页面 JSON 里那一列的真实内容）—— 逐行对账，顺序也一致
    payload = queries.fetch_table("sales", dimension=dimension, export=True,
                                  start="2011-11-01", end="2011-11-05")
    originals = [item["dimension_value"] for item in payload["items"]]
    assert [cell.value.date().isoformat() for cell in cells] == originals

    for cell in cells:
        assert isinstance(cell.value, (_dt.date, _dt.datetime)), \
            f"单元格不是日期类型：{cell.value!r}（{type(cell.value).__name__}）"
        assert is_date_format(cell.number_format), \
            f"number_format 不是合法日期格式：{cell.number_format!r}"
    # 缺陷本身：这一列里不许再出现 40875 那种"日期序列号被当数字显示"的单元格
    assert all(not isinstance(cell.value, int) for cell in cells)


def test_D1_期间列显示的就是日期而不是序列号(tmp_path):
    """缺陷的现象学断言：用户看到的是 `40875`（= 2011-11-28 的序列号），导出里不许再有它。"""
    body = _save({"source": "table", **SALES_QUERY}).json()
    _sheet, cells = _first_column_cells(_saved_path(body), "期间")
    values = [cell.value for cell in cells]
    assert not any(isinstance(value, int) for value in values), f"还有序列号：{values}"
    assert 40875 not in values
    assert all(isinstance(value, (_dt.date, _dt.datetime)) for value in values)
    assert {value.date().isoformat()[:4] for value in values} == {"2011"}


def test_D4_国家维度那一列仍是文本_没被当成日期转():
    """维度轴跟着维度走：按国家看时那是 `United Kingdom`，**不能**因为格式改动被转坏。"""
    from app.datasets import queries

    body = _save({"source": "table", **{**SALES_QUERY, "dimension": "country"}}).json()
    path = _saved_path(body)
    _sheet, cells = _first_column_cells(path, queries.SALES_AXIS_LABELS["country"])
    assert cells, "按国家维度导出没有数据"
    for cell in cells:
        assert isinstance(cell.value, str) and cell.value.strip(), cell.value
        assert not is_date_format(cell.number_format)


def test_D4_原始数据的下单时间仍是真时间():
    """别的导出位置也排查过：原始数据表的下单时间本来就是 date/datetime 列。"""
    body = _save({"source": "table", "table": "raw", "format": "xlsx",
                  "start": "2011-11-01", "end": "2011-11-02"}).json()
    path = _saved_path(body)
    _sheet, cells = _first_column_cells(path, "下单时间")
    assert cells, "原始数据没有导出任何行"
    for cell in cells:
        assert isinstance(cell.value, (_dt.date, _dt.datetime)), type(cell.value)
        assert is_date_format(cell.number_format)


def test_D4_客户表与商品表的日期列():
    """客户表的首次/最后购买本来就是 date 列；商品表里没有日期列。"""
    path = _saved_path(_save({"source": "table", "table": "customers", "format": "xlsx",
                              "start": "2011-11-01", "end": "2011-11-05"}).json())
    for label in ("首次购买", "最后购买"):
        _sheet, cells = _first_column_cells(path, label)
        assert cells
        for cell in cells:
            assert isinstance(cell.value, (_dt.date, _dt.datetime)), label
            assert is_date_format(cell.number_format), label


def test_D5_周报趋势表的日期没有回归():
    """D5：周报 xlsx 的「趋势」表本来就没有这个缺陷（日期以文本落盘，显示是对的）。

    这条把现状**钉住**：那一列读回来必须是可读的 `YYYY-MM-DD` 文本，
    绝不是序列号 —— 一旦有人把它也转成"日期值 + 文本格式"，这里立刻红。
    """
    chat = client.post("/api/chat", json={"question": WEEKLY_QUESTION, "use_llm": False})
    record = chat.json()
    report = record["report_document"]
    trend_table = report["sections"][2]["table"]

    body = _save({"source": "report", "conversation_id": record["conversation_id"],
                  "format": "xlsx"}).json()
    workbook = load_workbook(_saved_path(body))
    sheet = workbook["趋势"]
    rows = [[cell.value for cell in row] for row in sheet.iter_rows()]
    head = next(index for index, row in enumerate(rows)
                if row and str(row[0]).strip() == trend_table["headers"][0])
    data = [row for row in rows[head + 1:] if row[0] not in (None, "")]
    assert len(data) == len(trend_table["rows"]), "趋势表的行数变了"

    for row, original in zip(data, trend_table["rows"]):
        value = row[0]
        assert not isinstance(value, int), f"趋势表的日期成了序列号：{value}"
        assert str(value) == original[0], f"{value!r} != {original[0]!r}（趋势表日期读出来变了）"


def test_导出文件里没有日期值挂文本格式的病态组合(tmp_path):
    """把整份导出文件的每个单元格都扫一遍：**日期值 + `@` 文本格式** 的组合不许存在。

    这正是 `40875` 的成因（也是这次唯一改的地方）——只看「期间」一列会漏掉别处。
    """
    body = _save({"source": "table", **SALES_QUERY}).json()
    workbook = load_workbook(_saved_path(body))
    bad = []
    for sheet in workbook.worksheets:
        for row in sheet.iter_rows():
            for cell in row:
                if isinstance(cell.value, (_dt.date, _dt.datetime)) and cell.number_format == "@":
                    bad.append(f"{sheet.title}!{cell.coordinate}")
    assert not bad, f"这些单元格是'日期值 + 文本格式'（Excel 会显示成序列号）：{bad}"


def test_清洗函数不会把扩展名洗掉(tmp_path):
    """边界：名字里本来就有多个点时，只保留最后一段当扩展名。"""
    record = desktop.save_to_desktop("2011.11.28 销售表.xlsx", b"x")
    assert record["file_name"].endswith(".xlsx")
    assert os.path.exists(record["path"])
