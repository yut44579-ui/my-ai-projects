"""tests/test_renderer.py · TASK-002B 验收测试（Report Spec + Excel 渲染）。

════════════════════════════════════════════════════════════════════════
【本文件的验证思路 —— 为什么不是"自己写、自己读"就完事】
════════════════════════════════════════════════════════════════════════
被测：app/spec/models.py（Spec 模型）+ app/engine/renderer.py（渲染）

数字的真值只有一个来源：**executor 算出来的值**（002A 已用独立 Oracle 验过）。
本文件要证明的是"渲染没把数字改坏、没把样式弄丢"，所以用了三条**互相独立**的读取路径：
    路径 A  openpyxl 读回产出文件（用户/Excel 实际会走的路径）
    路径 B  **直接解压 xlsx 读原始 XML**（绕过 openpyxl 的读取逻辑，证明盘上存的就是这个数）
    路径 C  executor 的计算结果（金标准）
三条对不上就说明中间某一步出问题了。

样式部分（AC-03）不靠"肉眼看"，而是把**模板**和**产出**用 openpyxl 各读一遍逐属性比对
（字体名/字号/加粗/颜色、数字格式、列宽、合并区、行高）。

口径依据：docs/DECISIONS.md 的 D16（算数）与 D5（openpyxl 原地改，样式必须留住）。
"""

from __future__ import annotations

import hashlib
import json
import pathlib
import re
import sys
import zipfile

# 让 `import app` 在 `python -m pytest` / 直接 `pytest` 两种方式下都成立
PROJECT_ROOT = pathlib.Path(__file__).resolve().parents[1]
if str(PROJECT_ROOT) not in sys.path:
    sys.path.insert(0, str(PROJECT_ROOT))

import pytest  # noqa: E402
from openpyxl import load_workbook  # noqa: E402
from openpyxl.utils.cell import coordinate_to_tuple  # noqa: E402

from app.engine import executor, metrics, renderer  # noqa: E402
from app.spec import models as spec_models  # noqa: E402
from app.spec.models import ReportSpec  # noqa: E402

# 验收区间：2011-11-21 ~ 2011-11-27（周一~周日，数据覆盖内；与 002A 用同一周，便于交叉核对）
START = "2011-11-21"
END = "2011-11-27"

OUTPUTS_DIR = PROJECT_ROOT / "outputs"
TEMPLATE_PATH = PROJECT_ROOT / "templates" / "weekly_sales_template.xlsx"

# 产出文件名固定（不用时间戳），这样同一份测试跑两次能直接对比文件
RUN1_PATH = OUTPUTS_DIR / "pytest_weekly_sales_run1.xlsx"
RUN2_PATH = OUTPUTS_DIR / "pytest_weekly_sales_run2.xlsx"

# 读回容差：openpyxl 把数字写成 "%.16g"（openpyxl/compat/strings.py → cell/_writer.py），
# 双精度要 17 位有效数字才无损，故第 17 位可能差 1 个 ULP（相对 ≤ 2.3e-16）。
# 这里放到 1e-9：远大于 1 ULP（不误报），远小于 1 分钱（真写错了一定抓得住）。
READBACK_REL_TOL = 1e-9
READBACK_ABS_TOL = 1e-9

# 默认模板 6 个指标行（与 renderer.DEFAULT_METRIC_ROWS 对应）的 op。
# 写死在测试里（不 import renderer 的常量）：这里表达的是"我期望的口径"，不是"实现说是什么"。
EXPECTED_METRIC_OPS = {
    "sales_amount": "sum",
    "rows_in_range": "count",
    "rows_valid": "count",
    "rows_excluded": "count",
    "excluded_amount": "sum",
    "valid_qty_sum": "sum",
}


# ════════════════════════════════════════════════════════════════════════
# 打印小工具：让验收证据在一份输出里读得出来
# ════════════════════════════════════════════════════════════════════════
def banner(title: str) -> None:
    print("\n" + "─" * 72)
    print(title)
    print("─" * 72)


def spec_payload() -> dict:
    """一份完整的 Spec（JSON 形态），对应 renderer 生成的基础模板。

    故意**手写整份 dict**（不调任何"构造默认 Spec"的辅助函数）：
    测试要钉住的是"Spec 长这样、模板格子在那里"，实现改了这份期望就得跟着审视。
    """
    return {
        "spec_id": "weekly-sales",
        "name": "周销售报表",
        "version": 1,
        "data_source": {"kind": "excel", "path": "data/Online Retail.xlsx"},
        "time_range": {"start": START, "end": END},
        "metrics": [
            {"name": name, "dsl": {"op": op, "field": name}}
            for name, op in EXPECTED_METRIC_OPS.items()
        ],
        "group_by": [],
        "output": {
            "template": "templates/weekly_sales_template.xlsx",
            "sheet": renderer.DEFAULT_SHEET_NAME,
            "cells": dict(renderer.DEFAULT_CELL_MAP),
        },
        "created_at": "2026-09-23T10:00:00",
    }


# ════════════════════════════════════════════════════════════════════════
# fixtures（session 级：读 54 万行 xlsx 约百秒，算一次就够，全流程共用）
# ════════════════════════════════════════════════════════════════════════
@pytest.fixture(scope="session")
def weekly_spec() -> ReportSpec:
    return ReportSpec.from_dict(spec_payload())


@pytest.fixture(scope="session")
def execution_result() -> dict:
    """真算一遍（数字的唯一真值来源；这里不 mock、不写死）。"""
    return executor.compute_sales_amount(START, END)


@pytest.fixture(scope="session")
def values(weekly_spec: ReportSpec, execution_result: dict) -> dict:
    """executor 结果 → {指标名: 数值}（渲染的输入）。"""
    return renderer.values_from_result(weekly_spec, execution_result)


@pytest.fixture(scope="session")
def run1(weekly_spec: ReportSpec, values: dict) -> dict:
    """第 1 次渲染（固定路径，便于人工打开查看）。"""
    return renderer.render_report(weekly_spec, values, output_path=RUN1_PATH)


# ════════════════════════════════════════════════════════════════════════
# 辅助：三条独立读取路径
# ════════════════════════════════════════════════════════════════════════
def read_cells_via_openpyxl(path: pathlib.Path, sheet_name: str, addresses: list[str]) -> dict:
    """路径 A：openpyxl 读回（用户/Excel 实际走的路径）。"""
    workbook = load_workbook(path)
    sheet = workbook[sheet_name]
    return {address: sheet[address].value for address in addresses}


def read_cells_via_raw_xml(path: pathlib.Path, sheet_name: str, addresses: list[str]) -> dict:
    """路径 B：**绕过 openpyxl 读取**，直接解压 xlsx 读工作表 XML 里的 <v> 原始文本。

    xlsx 是 zip + xml，数字就写在 <c r="B4" t="n"><v>316412.16</v></c> 里。
    这条路径证明"盘上存的到底是什么"，不受 openpyxl 读取逻辑影响。
    工作表文件名：openpyxl 按 sheet 顺序写成 sheet1.xml / sheet2.xml …
    """
    workbook = load_workbook(path, read_only=True)          # 只为拿 sheet 顺序
    index = workbook.sheetnames.index(sheet_name)
    workbook.close()
    sheet_file = f"xl/worksheets/sheet{index + 1}.xml"
    with zipfile.ZipFile(path) as archive:
        xml = archive.read(sheet_file).decode("utf-8")

    found: dict[str, str] = {}
    for address in addresses:
        match = re.search(rf'<c r="{address}"[^>]*>\s*<v>([^<]*)</v>', xml)
        if match:
            found[address] = match.group(1)
    return found


def numeric_equal(actual: float, expected: float) -> tuple[bool, float, float]:
    """数字等价判断（含 openpyxl 序列化容差），返回 (是否相等, 绝对差, 容差)。"""
    abs_diff = abs(float(actual) - float(expected))
    tolerance = max(READBACK_ABS_TOL, abs(float(expected)) * READBACK_REL_TOL)
    return abs_diff <= tolerance, abs_diff, tolerance


def style_snapshot(path: pathlib.Path, sheet_name: str) -> dict:
    """把"必须留住"的样式属性抓成一份快照，用于模板 vs 产出逐项比对（AC-03）。"""
    workbook = load_workbook(path)
    sheet = workbook[sheet_name]
    data_rows = range(
        renderer.DEFAULT_FIRST_DATA_ROW,
        renderer.DEFAULT_FIRST_DATA_ROW + len(renderer.DEFAULT_METRIC_ROWS),
    )
    return {
        # 字体：标题（名称/字号/加粗/颜色）、表头（加粗/颜色）、数据格（字体名）
        "title_font": {
            "name": sheet["A1"].font.name,
            "size": sheet["A1"].font.size,
            "bold": sheet["A1"].font.bold,
            "color": sheet["A1"].font.color.rgb if sheet["A1"].font.color else None,
        },
        "title_alignment": sheet["A1"].alignment.horizontal,
        "title_fill": sheet["A1"].fill.fgColor.rgb,
        "header_font": {
            "bold": sheet["A3"].font.bold,
            "color": sheet["A3"].font.color.rgb if sheet["A3"].font.color else None,
        },
        "header_fill": sheet["A3"].fill.fgColor.rgb,
        # 数字格式：每个数据格
        "number_formats": {f"B{row}": sheet[f"B{row}"].number_format for row in data_rows},
        # 列宽
        "column_widths": {
            column: sheet.column_dimensions[column].width for column in ("A", "B", "C")
        },
        # 合并单元格 + 行高 + 标签文案
        "merged_ranges": sorted(str(rng) for rng in sheet.merged_cells.ranges),
        "row_heights": {row: sheet.row_dimensions[row].height for row in (1, 2, 3)},
        "labels": {f"A{row}": sheet[f"A{row}"].value for row in data_rows},
        "row3_headers": [sheet.cell(row=3, column=col).value for col in (1, 2, 3)],
    }


# ════════════════════════════════════════════════════════════════════════
# AC-01（前半）· Spec 模型：JSON 往返 / 版本递增 / 非法输入一律报错
# ════════════════════════════════════════════════════════════════════════
def test_ac01_spec_model_roundtrip_and_strictness() -> None:
    banner("AC-01 Spec 模型：to_dict/from_dict 往返 + version 递增 + 严格校验")
    spec = ReportSpec.from_dict(spec_payload())

    # ① to_dict 的产物必须是**纯 JSON**（D11：Spec 要能落库、能审计）
    payload = spec.to_dict()
    text = json.dumps(payload, ensure_ascii=False, sort_keys=True)
    print(f"① to_dict 可 json.dumps：{len(text)} 字符，前 100 字符：{text[:100]}…")
    assert isinstance(payload["time_range"]["start"], str), "日期必须序列化成字符串（JSON 友好）"
    assert payload["time_range"]["start"] == START

    # ② 往返相等：from_dict(to_dict()) == 原对象
    restored = ReportSpec.from_dict(json.loads(text))
    print(f"② 往返相等：{restored == spec}")
    assert restored == spec, "JSON 往返后 Spec 必须完全相等"

    # ③ version 递增：原对象冻结不变，新对象 version+1（D11 版本冻结）
    bumped = spec.bump_version()
    print(f"③ bump_version：{spec.version} → {bumped.version}；原对象 version 仍为 {spec.version}")
    assert bumped.version == spec.version + 1
    assert spec.version == 1, "已确认的 Spec 不能被原地改（frozen）"
    assert bumped.to_dict() | {"version": 1} == spec.to_dict(), "除 version 外其余字段必须逐字不变"
    with pytest.raises(Exception):
        spec.version = 99            # frozen：不许原地改口径

    # ④ 口径默认值来自 metrics.py（单一来源，避免两处口径漂移 —— D12）
    print(f"④ 默认时间字段={spec.time_range.time_field!r}，"
          f"区间语义含'含首尾全天'={'含首尾全天' in spec.time_range.semantics}，"
          f"默认排除={spec.exclusion_keys()}")
    assert spec.time_range.time_field == metrics.TIME_FIELD
    assert spec.time_range.semantics == metrics.RANGE_SEMANTICS
    assert spec.exclusion_keys() == tuple(rule.key for rule in metrics.EXCLUSION_RULES)
    assert spec_models.DEFAULT_EXCLUSIONS == tuple(rule.key for rule in metrics.EXCLUSION_RULES)

    # ⑤ 受限 DSL 的算子白名单与"排除规则键"都必须与 metrics.py 对得上（防漂移）
    from typing import get_args

    literal_exclusions = set(get_args(spec_models.ExclusionKey))
    rule_keys = {rule.key for rule in metrics.EXCLUSION_RULES}
    print(f"⑤ ExclusionKey 白名单={sorted(literal_exclusions)}；metrics 规则键={sorted(rule_keys)}")
    assert literal_exclusions == rule_keys, "Spec 的排除规则白名单与 metrics.py 规则表不一致"

    # ⑥ 非法输入必须报错（不静默、不猜 —— 铁律 4）
    cases: list[tuple[str, dict]] = []
    base = spec_payload()

    def mutated(change) -> dict:
        clone = json.loads(json.dumps(base))
        change(clone)
        return clone

    cases.append(("LLM 多输出一个字段 extra", mutated(lambda c: c.update(surprise=1))))
    cases.append(("算子不在白名单（median）", mutated(lambda c: c["metrics"][0]["dsl"].update(op="median"))))
    cases.append(("sum 缺 field", mutated(lambda c: c["metrics"][0]["dsl"].pop("field"))))
    cases.append(("yoy 缺 of_metric", mutated(lambda c: c["metrics"][0]["dsl"].update(op="yoy", field=None))))
    cases.append(("of_metric 指向不存在的指标", mutated(
        lambda c: c["metrics"][0]["dsl"].update(op="yoy", field=None, of_metric="nope"))))
    cases.append(("cells 里写了未声明的指标名", mutated(lambda c: c["output"]["cells"].update(sale_amount="B4"))))
    cases.append(("单元格地址小写", mutated(lambda c: c["output"]["cells"].update(sales_amount="b4"))))
    cases.append(("单元格地址超过 Excel 列上限", mutated(lambda c: c["output"]["cells"].update(sales_amount="ZZZZ4"))))
    cases.append(("exclude 出现未知规则名", mutated(lambda c: c["metrics"][0]["dsl"].update(exclude=["whatever"]))))
    cases.append(("同 Spec 内 exclude 口径不一致", mutated(lambda c: c["metrics"][1]["dsl"].update(exclude=["cancelled"]))))
    cases.append(("指标名重复", mutated(lambda c: c["metrics"][1].update(name="sales_amount"))))
    cases.append(("start 晚于 end", mutated(lambda c: c["time_range"].update(start="2012-01-01"))))
    cases.append(("template 不是 xlsx", mutated(lambda c: c["output"].update(template="templates/t.xls"))))
    cases.append(("excel 数据源缺 path", mutated(lambda c: c["data_source"].pop("path"))))

    for label, payload_bad in cases:
        with pytest.raises(Exception) as excinfo:
            ReportSpec.from_dict(payload_bad)
        print(f"   ⑥ 拒绝[{label}] → {type(excinfo.value).__name__}")
        assert "ValidationError" in type(excinfo.value).__name__ or isinstance(excinfo.value, ValueError)


# ════════════════════════════════════════════════════════════════════════
# AC-02 · 产出文件读回来 == executor 的计算值（三条独立路径互相印证）
# ════════════════════════════════════════════════════════════════════════
def test_ac02_readback_matches_executor(
    weekly_spec: ReportSpec, execution_result: dict, values: dict, run1: dict
) -> None:
    banner("AC-02 产出 xlsx 读回 == executor 计算值（openpyxl 路径 + 原始 XML 路径）")
    cell_map = weekly_spec.cell_map()
    addresses = list(cell_map.values())

    via_openpyxl = read_cells_via_openpyxl(pathlib.Path(run1["output_path"]), weekly_spec.output.sheet, addresses)
    via_xml = read_cells_via_raw_xml(pathlib.Path(run1["output_path"]), weekly_spec.output.sheet, addresses)

    print(f"executor 原始结果：amount={execution_result['amount']!r}, "
          f"rows_in_range={execution_result['rows_in_range']}, rows_excluded={execution_result['rows_excluded']}, "
          f"valid_qty_sum={execution_result['valid_qty_sum']!r}")
    print(f"{'指标':<18}{'格':<5}{'executor':<24}{'openpyxl 读回':<24}{'差值':<14}{'XML 原文'}")
    for metric_name, address in cell_map.items():
        expected = values[metric_name]
        read_back = via_openpyxl[address]
        ok, abs_diff, tolerance = numeric_equal(read_back, expected)
        print(f"{metric_name:<18}{address:<5}{expected!r:<24}{read_back!r:<24}{abs_diff!r:<14}{via_xml.get(address)!r}")
        assert ok, f"AC-02 FAIL：{address}（{metric_name}）读回 {read_back!r} != executor {expected!r}"
        # 路径 A 与路径 B 必须一致（openpyxl 没读错、盘上就是这个数）
        assert float(via_xml[address]) == float(read_back), f"AC-02 FAIL：XML 原文与 openpyxl 读回不一致（{address}）"

    # ── 整数指标必须**严格相等**（行数/计数不该有任何误差）─────────────────
    for metric_name in ("rows_in_range", "rows_valid", "rows_excluded"):
        assert via_openpyxl[cell_map[metric_name]] == values[metric_name], f"{metric_name} 必须严格相等"
    print("整数指标（区间行数/有效行数/被排除行数）严格相等 ✓")

    # ── 浮点指标的差异必须能解释成"openpyxl 的序列化格式"，不是我方算错 ────
    # openpyxl 用 "%.16g" 写数字（见 compat/strings.py 的 safe_string），第 17 位会丢。
    # 断言盘上的 XML 原文恰好等于 "%.16g" % 真值 —— 把"损耗来自序列化"钉成可检验的事实。
    amount_addr = cell_map["sales_amount"]
    expected_amount = execution_result["amount"]
    print(f"浮点指标 {amount_addr}：XML 原文={via_xml[amount_addr]!r}，"
          f"'%.16g' % 真值={'%.16g' % expected_amount!r}，真值 repr={expected_amount!r}")
    assert via_xml[amount_addr] == "%.16g" % expected_amount, "AC-02 FAIL：XML 里的数字不是 openpyxl 的 %.16g 序列化结果"
    assert float(via_xml[amount_addr]) == pytest.approx(
        expected_amount, rel=READBACK_REL_TOL, abs=READBACK_ABS_TOL
    )

    # ── 渲染输入必须是 executor 的**原值**（渲染环节不加工数字）───────────
    for metric_name, expected in (
        ("sales_amount", execution_result["amount"]),
        ("rows_in_range", execution_result["rows_in_range"]),
        ("rows_valid", execution_result["rows_valid"]),
        ("rows_excluded", execution_result["rows_excluded"]),
        ("excluded_amount", execution_result["excluded_amount"]),
        ("valid_qty_sum", execution_result["valid_qty_sum"]),
    ):
        assert values[metric_name] == expected, f"values[{metric_name}] 与 executor 原值不等"
    print("values_from_result 抽取的值与 executor 原值**逐一严格相等**（渲染不二次加工）✓")

    # ── 渲染器自己的读回自检也必须通过（产品内置的 D10 守卫）─────────────
    verification = run1["verification"]
    print(f"渲染器内置读回自检：passed={verification['passed']}，"
          f"容差 rel={verification['rel_tol']} abs={verification['abs_tol']}")
    assert verification["passed"] is True
    assert all(item["passed"] for item in verification["items"])
    assert all(item["style_passed"] for item in verification["items"]), "数字格式在渲染后必须保留"


# ════════════════════════════════════════════════════════════════════════
# AC-03 · 模板样式保留（字体 / 数字格式 / 列宽 … 模板 vs 产出逐项比对）
# ════════════════════════════════════════════════════════════════════════
def test_ac03_template_styles_preserved(weekly_spec: ReportSpec, run1: dict) -> None:
    banner("AC-03 样式保留：模板 vs 产出（字体 / 数字格式 / 列宽 / 合并区 / 行高）")
    assert TEMPLATE_PATH.exists(), f"模板必须先存在：{TEMPLATE_PATH}"
    template_style = style_snapshot(TEMPLATE_PATH, weekly_spec.output.sheet)
    output_style = style_snapshot(pathlib.Path(run1["output_path"]), weekly_spec.output.sheet)

    print(f"{'属性':<26}{'模板':<44}{'产出':<44}一致")
    checked = 0
    for key in template_style:
        before, after = template_style[key], output_style[key]
        same = before == after
        checked += 1
        print(f"{key:<26}{str(before)[:42]:<44}{str(after)[:42]:<44}{'✓' if same else '✗'}")
        assert same, f"AC-03 FAIL：样式属性 {key} 在渲染后变了：模板={before!r} 产出={after!r}"

    # 逐项把"至少各一项"钉死（AC-03 明确要求字体 / 数字格式 / 列宽各至少一项）
    assert output_style["title_font"]["name"] == template_style["title_font"]["name"]
    assert output_style["title_font"]["size"] == template_style["title_font"]["size"]
    assert output_style["title_font"]["bold"] is True, "标题必须是加粗的（模板样式）"
    assert output_style["header_font"]["bold"] is True, "表头必须加粗"
    assert output_style["number_formats"]["B4"] == "#,##0.00", "金额格必须保留两位小数格式"
    assert output_style["number_formats"]["B5"] == "#,##0", "行数格必须是整数格式"
    assert output_style["column_widths"]["A"] == template_style["column_widths"]["A"]
    assert output_style["merged_ranges"] == template_style["merged_ranges"], "合并标题必须保留"
    assert output_style["labels"] == template_style["labels"], "指标标签文案不能被渲染改动"
    assert output_style["row3_headers"] == template_style["row3_headers"]
    print(f"逐项断言：字体(名称/字号/加粗) ✓  数字格式({output_style['number_formats']['B4']}) ✓  "
          f"列宽(A={output_style['column_widths']['A']}) ✓  合并区{output_style['merged_ranges']} ✓")

    # 数据格写完后**标签/说明文字**必须原样（渲染只改数据格，不动别的）
    workbook = load_workbook(pathlib.Path(run1["output_path"]))
    sheet = workbook[weekly_spec.output.sheet]
    print(f"标题文字保持：{sheet['A1'].value!r}；副标题保持：{sheet['A2'].value!r}")
    assert sheet["A1"].value == "周销售报表"
    assert "含首尾全天" in sheet["A2"].value, "副标题（口径说明）必须还在"
    assert checked >= 8, "样式比对项太少，AC-03 没有真正逐项验证"


# ════════════════════════════════════════════════════════════════════════
# AC-04 · 可重复性：同一 Spec 跑两次，数字部分完全一致
# ════════════════════════════════════════════════════════════════════════
def test_ac04_repeatable_numbers(weekly_spec: ReportSpec, values: dict, run1: dict) -> None:
    banner("AC-04 可重复性：同一 Spec 渲染两次，数字部分必须完全一致")
    run2 = renderer.render_report(weekly_spec, values, output_path=RUN2_PATH)
    print(f"第 1 次：{run1['output_path']}（{run1['size_bytes']} bytes）")
    print(f"第 2 次：{run2['output_path']}（{run2['size_bytes']} bytes）")

    cells = list(weekly_spec.cell_map().values())
    cells1 = read_cells_via_openpyxl(pathlib.Path(run1["output_path"]), weekly_spec.output.sheet, cells)
    cells2 = read_cells_via_openpyxl(pathlib.Path(run2["output_path"]), weekly_spec.output.sheet, cells)
    cells1_xml = read_cells_via_raw_xml(pathlib.Path(run1["output_path"]), weekly_spec.output.sheet, cells)
    cells2_xml = read_cells_via_raw_xml(pathlib.Path(run2["output_path"]), weekly_spec.output.sheet, cells)

    for address in cells:
        print(f"  {address}: 第一次={cells1[address]!r}  第二次={cells2[address]!r}")
        assert cells1[address] == cells2[address], f"AC-04 FAIL：{address} 两次渲染结果不同"
        assert cells1_xml[address] == cells2_xml[address], f"AC-04 FAIL：{address} 两次盘上 XML 不同"

    # 渲染输入也一致（Spec → values → 文件，整条链可重复）
    assert run1["cells_written"] == run2["cells_written"]

    # 说明：两次产出的**文件字节**可能不同（xlsx 里有创建/修改时间戳元数据），
    # 但 AC-04 要求的是"数字部分完全一致"，上面已逐格验证。
    bytes_equal = pathlib.Path(run1["output_path"]).read_bytes() == pathlib.Path(run2["output_path"]).read_bytes()
    print(f"整个文件字节相同：{bytes_equal}（不影响 AC-04 —— 数字部分已逐格比对一致）")
    print(f"  run1 sha256={hashlib.sha256(pathlib.Path(run1['output_path']).read_bytes()).hexdigest()[:16]}…")
    print(f"  run2 sha256={hashlib.sha256(pathlib.Path(run2['output_path']).read_bytes()).hexdigest()[:16]}…")


# ════════════════════════════════════════════════════════════════════════
# AC-05 · 文件落到 outputs/（默认路径），打印路径 + 大小
# ════════════════════════════════════════════════════════════════════════
def test_ac05_output_lands_in_outputs_dir(weekly_spec: ReportSpec, values: dict) -> None:
    banner("AC-05 默认输出路径：必须落在 outputs/ 下")
    result = renderer.render_report(weekly_spec, values)          # 不指定 output_path → 走默认
    path = pathlib.Path(result["output_path"])
    print(f"产出路径 : {path}")
    print(f"文件大小 : {result['size_bytes']} bytes")
    print(f"模板来源 : {result['template']}")
    print(f"耗时     : {result['seconds']:.3f}s；sheet={result['sheet']!r}；spec={result['spec_id']} v{result['spec_version']}")
    print(f"单元格   : " + ", ".join(f"{info['cell']}={info['value']!r}" for info in result["cells_written"].values()))

    assert path.exists(), "AC-05 FAIL：产出文件不存在"
    assert path.parent == OUTPUTS_DIR, f"AC-05 FAIL：产出没落在 outputs/（实际 {path.parent}）"
    assert path.suffix == ".xlsx"
    assert result["size_bytes"] > 0
    assert path.stat().st_size == result["size_bytes"]
    # 模板不能被就地改动（渲染是"模板副本 → 产出"，D5）
    assert path.resolve() != TEMPLATE_PATH.resolve()


# ════════════════════════════════════════════════════════════════════════
# 基础模板：必须由代码生成、结构完整、且**不是** pandas 重建的
# ════════════════════════════════════════════════════════════════════════
def test_template_generated_by_code(tmp_path: pathlib.Path) -> None:
    banner("模板：由 openpyxl 代码生成（D5 禁止 pandas.to_excel 重建），结构完整")
    target = tmp_path / "generated_template.xlsx"
    built = renderer.build_default_template(target)
    print(f"生成模板：{built}（{built.stat().st_size} bytes）")
    assert built.exists()

    # ① 已存在时不许静默覆盖（防手滑覆盖用户自定义模板）
    with pytest.raises(renderer.RenderError):
        renderer.build_default_template(target)
    print("① 已存在时拒绝覆盖（需显式 overwrite=True）✓")

    # ② 结构完整：标题合并 / 表头 / 数据格留空 / 数字格式 / 列宽
    snapshot = style_snapshot(built, renderer.DEFAULT_SHEET_NAME)
    print(f"② 合并区={snapshot['merged_ranges']}；数字格式={snapshot['number_formats']}")
    print(f"   列宽={snapshot['column_widths']}；标签={snapshot['labels']}")
    assert snapshot["title_font"]["bold"] is True and snapshot["title_font"]["size"] >= 14
    assert snapshot["merged_ranges"], "标题必须是合并单元格（pandas 重建会丢这个）"
    assert snapshot["number_formats"]["B4"] == "#,##0.00"
    assert snapshot["column_widths"]["A"] == 26.0 and snapshot["column_widths"]["C"] == 62.0
    workbook = load_workbook(built)
    sheet = workbook[renderer.DEFAULT_SHEET_NAME]
    for address in renderer.DEFAULT_CELL_MAP.values():
        assert sheet[address].value is None, f"模板里的数据格必须是空的（{address} 有值）"
    print(f"   数据格 {sorted(renderer.DEFAULT_CELL_MAP.values())} 全部为空（等渲染填）✓")

    # ③ 默认模板与 DEFAULT_CELL_MAP / DEFAULT_METRIC_ROWS 必须对得上（防两处漂移）
    rows = range(
        renderer.DEFAULT_FIRST_DATA_ROW,
        renderer.DEFAULT_FIRST_DATA_ROW + len(renderer.DEFAULT_METRIC_ROWS),
    )
    for (name, label, _fmt, _note), row in zip(renderer.DEFAULT_METRIC_ROWS, rows):
        assert renderer.DEFAULT_CELL_MAP[name] == f"B{row}"
        assert sheet[f"A{row}"].value == label
    print(f"③ DEFAULT_METRIC_ROWS ↔ DEFAULT_CELL_MAP ↔ 模板三者行序一致（{len(rows)} 行）✓")

    # ④ 渲染模块不得**调用** pandas.to_excel（D5 明确否掉"重建文件"；注释里提到它不算违规）
    source = (PROJECT_ROOT / "app" / "engine" / "renderer.py").read_text(encoding="utf-8")
    assert "to_excel(" not in source, "renderer.py 调用了 to_excel —— 违反 D5（必须原地改，不许重建）"
    assert "import pandas" not in source, "renderer.py 不该 import pandas（渲染与算数分层）"
    print("④ renderer.py 无 to_excel( 调用 / 无 pandas 依赖 ✓")


# ════════════════════════════════════════════════════════════════════════
# Spec ↔ executor 的桥接：口径固化真的能驱动计算
# ════════════════════════════════════════════════════════════════════════
def test_executor_kwargs_follow_spec_exclusions(execution_result: dict) -> None:
    banner("桥接：Spec 里固化的排除规则 → executor 的开关参数（口径由 Spec 决定）")

    # ① 默认口径（三条全开）→ 开关全 True，且与 002A 默认值一致
    default_spec = ReportSpec.from_dict(spec_payload())
    default_kwargs = renderer.executor_kwargs(default_spec)
    print(f"① 默认口径 kwargs={default_kwargs}")
    assert default_kwargs == {
        "exclude_cancelled": True, "exclude_negative_qty": True, "exclude_nonpositive_price": True,
    }

    # ② 关掉"取消单"规则 → 参数名/含义由 metrics.py 的规则表决定（不手写字符串）
    payload = spec_payload()
    for metric in payload["metrics"]:
        metric["dsl"]["exclude"] = ["negative_qty", "nonpositive_price"]
    spec_off = ReportSpec.from_dict(payload)
    kwargs_off = renderer.executor_kwargs(spec_off)
    print(f"② 只排除退货/负单价 kwargs={kwargs_off}")
    assert kwargs_off == {
        "exclude_cancelled": False, "exclude_negative_qty": True, "exclude_nonpositive_price": True,
    }

    # ③ 用 Spec 的口径真算一遍：统计口径必须跟着 Spec 走
    result_off = executor.compute_sales_amount(START, END, **kwargs_off)
    print(f"③ 默认口径：金额={execution_result['amount']!r} 被排除={execution_result['rows_excluded']} "
          f"明细={execution_result['excluded_detail']}")
    print(f"   关掉取消单：金额={result_off['amount']!r} 被排除={result_off['rows_excluded']} "
          f"明细={result_off['excluded_detail']}")
    assert result_off["excluded_detail"]["cancelled"] == 0, "被关闭的规则计数必须为 0（口径只认已启用规则）"
    assert result_off["rows_in_range"] == execution_result["rows_in_range"], "第 1 步时间筛选不受开关影响"
    assert result_off["rows_excluded"] <= execution_result["rows_excluded"], "少一条规则，被排除行数只能变少或不变"
    # 注意（别把它当成 bug）：本数据集里取消单行**通常同时是负数量行**（Cancellation 的 Quantity<0），
    # 关掉 cancelled 后这些行仍被 negative_qty 规则排除，所以金额可能**不变** —— 这是口径的正确表现。
    # "开关真的接到了计算上"由下面 ④ 证明：三条全关时金额必须变化。
    if result_off["amount"] == execution_result["amount"]:
        print("   金额未变：取消单行同时命中「负数量」规则，仍被排除（口径正确，不是开关失效）")

    # ④ 三条全关（Spec 里 exclude=[]）→ 金额必须变化（否则说明 Spec 的口径根本没驱动计算）
    payload_all_off = spec_payload()
    for metric in payload_all_off["metrics"]:
        metric["dsl"]["exclude"] = []
    kwargs_all_off = renderer.executor_kwargs(ReportSpec.from_dict(payload_all_off))
    result_all_off = executor.compute_sales_amount(START, END, **kwargs_all_off)
    print(f"④ 三条全关 kwargs={kwargs_all_off} → 金额={result_all_off['amount']!r}"
          f"（默认口径 {execution_result['amount']!r}）")
    assert kwargs_all_off == {
        "exclude_cancelled": False, "exclude_negative_qty": False, "exclude_nonpositive_price": False,
    }
    assert result_all_off["amount"] != execution_result["amount"], "口径全关后金额必须变化（Spec 口径没接上计算）"
    assert result_all_off["rows_excluded"] == 0 and result_all_off["excluded_detail"]["multi_rule_hit"] == 0
    print(f"   全关后被排除行数={result_all_off['rows_excluded']}（应为 0），"
          f"区间全部行金额={result_all_off['amount']!r}")

    # ⑤ 参数名必须来自 metrics.EXCLUSION_RULES[].param（防自造参数名）
    params = {rule.param for rule in metrics.EXCLUSION_RULES}
    print(f"⑤ kwargs 键={sorted(default_kwargs)}；metrics 参数名={sorted(params)}")
    assert set(default_kwargs) == params


# ════════════════════════════════════════════════════════════════════════
# 错误路径：全都必须"报错 + 不产出文件"（D10）
# ════════════════════════════════════════════════════════════════════════
def test_error_paths_fail_loudly(weekly_spec: ReportSpec, values: dict, tmp_path: pathlib.Path) -> None:
    banner("错误路径：模板缺失 / sheet 不存在 / 格子越界 / 合并区 / 缺数值 / group_by / 产不出的口径")
    sheet_name = weekly_spec.output.sheet

    # ① 模板缺失 → 报错，且错误信息里给出可选项（不猜、不静默换模板）
    with pytest.raises(renderer.RenderError) as exc:
        payload = spec_payload()
        payload["output"]["template"] = "templates/不存在的模板.xlsx"
        spec_missing = ReportSpec.from_dict(payload)
        renderer.render_report(spec_missing, values, output_path=tmp_path / "x.xlsx")
    print(f"① 模板缺失：{str(exc.value).splitlines()[0]}")
    assert "不存在" in str(exc.value) and "build_default_template" in str(exc.value)
    assert not (tmp_path / "x.xlsx").exists(), "失败时不许留下文件"

    # ② 传 create_template_if_missing=True → 按 TASK-002 第 9 条用代码生成模板并渲染成功
    payload = spec_payload()
    target_template = tmp_path / "auto_template.xlsx"
    payload["output"]["template"] = str(target_template)
    spec_auto = ReportSpec.from_dict(payload)
    auto_out = renderer.render_report(
        spec_auto, values, output_path=tmp_path / "auto_out.xlsx", create_template_if_missing=True
    )
    print(f"② create_template_if_missing=True → 自动生成模板并产出：{auto_out['output_path']}")
    assert target_template.exists() and pathlib.Path(auto_out["output_path"]).exists()

    # ③ sheet 不存在 → 错误信息里列出可选的 sheet（铁律 4：列候选让用户选）
    with pytest.raises(renderer.RenderError) as exc:
        payload = spec_payload()
        payload["output"]["sheet"] = "不存在的表"
        renderer.render_report(ReportSpec.from_dict(payload), values, output_path=tmp_path / "y.xlsx")
    print(f"③ sheet 不存在：{str(exc.value).splitlines()[-1].strip()}")
    assert sheet_name in str(exc.value), "错误信息必须列出模板里真实存在的 sheet"

    # ④ 数据格超出模板已用范围（地址打错）→ 报错
    with pytest.raises(renderer.RenderError) as exc:
        payload = spec_payload()
        payload["output"]["cells"]["sales_amount"] = "B40"
        renderer.render_report(ReportSpec.from_dict(payload), values, output_path=tmp_path / "z.xlsx")
    print(f"④ 格子越界：{str(exc.value).splitlines()[0]}")
    assert "超出模板已用范围" in str(exc.value)
    assert not (tmp_path / "z.xlsx").exists()

    # ⑤ 写到合并区（非左上角）→ 报错（否则 Excel 会静默忽略这个值）
    with pytest.raises(renderer.RenderError) as exc:
        payload = spec_payload()
        payload["output"]["cells"]["sales_amount"] = "B1"     # A1:C1 是合并标题
        renderer.render_report(ReportSpec.from_dict(payload), values, output_path=tmp_path / "w.xlsx")
    print(f"⑤ 合并区：{str(exc.value).splitlines()[0]}")
    assert "合并区" in str(exc.value)

    # ⑥ 缺某个指标的数值 → 报错，并列出缺了谁（绝不填 0）
    with pytest.raises(renderer.RenderError) as exc:
        renderer.render_report(weekly_spec, {"sales_amount": 1.0}, output_path=tmp_path / "v.xlsx")
    print(f"⑥ 缺数值：{str(exc.value).splitlines()[0]}")
    assert "缺少这些指标的数值" in str(exc.value)
    assert not (tmp_path / "v.xlsx").exists()

    # ⑦ group_by 非空 → 明细表渲染口径未定义，明确拒绝而不是画一张错的表
    with pytest.raises(renderer.RenderError) as exc:
        payload = spec_payload()
        payload["group_by"] = ["Country"]
        renderer.render_report(ReportSpec.from_dict(payload), values, output_path=tmp_path / "u.xlsx")
    print(f"⑦ group_by：{str(exc.value).splitlines()[0]}")
    assert "group_by" in str(exc.value)

    # ⑧ 值不是数字（把 range 这种列表塞进数据格）→ 报错
    with pytest.raises(renderer.RenderError):
        renderer.render_report(weekly_spec, values | {"sales_amount": [1, 2]}, output_path=tmp_path / "t.xlsx")
    with pytest.raises(renderer.RenderError):
        renderer.render_report(weekly_spec, values | {"sales_amount": True}, output_path=tmp_path / "s.xlsx")
    print("⑧ 非数字（list / bool）被拒 ✓")


def test_unsupported_metric_dsl_rejected(weekly_spec: ReportSpec, execution_result: dict) -> None:
    banner("受限 DSL 里合法但当前引擎产不出的口径 → 明确报错，并列出可用组合")
    payload = spec_payload()
    payload["metrics"].append({"name": "avg_qty", "dsl": {"op": "mean", "field": "Quantity"}})
    payload["metrics"].append({"name": "sales_yoy", "dsl": {"op": "yoy", "of_metric": "sales_amount"}})
    spec = ReportSpec.from_dict(payload)          # Spec 层面合法（受限 DSL 允许这些算子）
    print(f"Spec 校验通过（算子白名单内的合法结构）：{[m.name for m in spec.metrics]}")

    with pytest.raises(renderer.UnsupportedMetricError) as exc:
        renderer.values_from_result(spec, execution_result)
    message = str(exc.value)
    print("报错信息：")
    for line in message.splitlines():
        print("   " + line)
    assert "avg_qty" in message
    assert "sum(sales_amount)" in message, "报错必须列出当前支持的组合（列候选，不猜）"
    assert issubclass(renderer.UnsupportedMetricError, renderer.RenderError)


def test_renderer_does_not_recompute_numbers(weekly_spec: ReportSpec, execution_result: dict) -> None:
    banner("边界：渲染层不加工数字 —— 传什么写什么，真实性由 executor 保证")
    values = renderer.values_from_result(weekly_spec, execution_result)
    # 用一份**明显不同**的 values 渲染：文件里必须是传进去的值（说明渲染没有偷偷自己算）
    probe = {name: value * 10 for name, value in values.items()}
    result = renderer.render_report(weekly_spec, probe, output_path=OUTPUTS_DIR / "pytest_probe.xlsx")
    read_back = read_cells_via_openpyxl(
        pathlib.Path(result["output_path"]), weekly_spec.output.sheet, list(weekly_spec.cell_map().values())
    )
    for metric_name, address in weekly_spec.cell_map().items():
        ok, abs_diff, _tol = numeric_equal(read_back[address], probe[metric_name])
        print(f"  {metric_name:<18}{address:<5}写入={probe[metric_name]!r:<24}读回={read_back[address]!r:<24}差={abs_diff!r}")
        assert ok, f"渲染层改了数字：{address} 写 {probe[metric_name]!r} 读回 {read_back[address]!r}"
    pathlib.Path(result["output_path"]).unlink()          # 探针文件删掉，不留在 outputs/
    print("探针文件已删除（outputs/ 只留正式产出）✓")


def test_cell_address_helpers() -> None:
    banner("单元格地址工具：合法/非法边界")
    assert spec_models.column_index("A") == 1
    assert spec_models.column_index("Z") == 26
    assert spec_models.column_index("AA") == 27
    assert spec_models.column_index("XFD") == 16384
    print("column_index: A=1, Z=26, AA=27, XFD=16384 ✓")
    for good in ("A1", "B4", "AA100", "XFD1048576"):
        assert spec_models.validate_cell_address(good) == good
    print("合法地址通过：A1 / B4 / AA100 / XFD1048576")
    for bad in ("b4", "4B", "B0", "A1:B2", "", "XFE1", "A1048577"):
        with pytest.raises(ValueError):
            spec_models.validate_cell_address(bad)
        print(f"非法地址被拒：{bad!r}")
    # coordinate_to_tuple 与 column_index 必须一致（两套实现不能打架）
    for address in ("A1", "B4", "AA100"):
        row, column = coordinate_to_tuple(address)
        letters = re.match(r"^([A-Z]+)", address).group(1)
        assert column == spec_models.column_index(letters) and row == int(address[len(letters):])
    print("column_index 与 openpyxl 的 coordinate_to_tuple 结果一致 ✓")
