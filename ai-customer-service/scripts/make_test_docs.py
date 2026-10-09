"""造真实的 Word / PDF / Excel 测试文件。

★ 为什么必须用**真实的二进制文件**测，而不是 mock 解析器：
    解析器最容易出的错是"格式理解错了"（表格顺序乱、合并单元格丢了、
    扫描件读出一堆乱码）。mock 掉解析器就永远测不到这些。

    而且客户的实际资料就是这些格式 —— 用 markdown 测出来的"通过"
    不代表接了真实资料也能用。
"""

from __future__ import annotations

import sys
from pathlib import Path

try:
    sys.stdout.reconfigure(encoding="utf-8", errors="replace")
except Exception:
    pass

OUT = Path(sys.argv[1] if len(sys.argv) > 1 else r"D:\ai-kefu\data\kb_real")
OUT.mkdir(parents=True, exist_ok=True)


def make_docx() -> Path:
    """Word：含标题、正文、表格。

    ★ 表格是重点：企业内部制度里大量信息在表格里
      （出差标准、审批权限、岗位对照），
      只读段落不读表格的话，索引里就少了关键的一半。
    """
    import docx

    d = docx.Document()
    d.add_heading("费用报销管理办法", level=1)
    d.add_paragraph("标题: 费用报销管理办法")
    d.add_paragraph("权限: external")
    d.add_paragraph("---")
    d.add_paragraph("第3.2条 报销单经审批通过后，应于5个工作日内完成付款。")
    d.add_paragraph("第3.3条 发票必须贴在报销单背面，拍照上传后提交。")
    d.add_paragraph("以下是差旅费标准：")
    t = d.add_table(rows=3, cols=3)
    hdr = ["城市级别", "住宿费上限", "交通"]
    for i, h in enumerate(hdr):
        t.cell(0, i).text = h
    rows = [
        ["一线城市", "500 元/晚", "实报实销"],
        ["其他城市", "350 元/晚", "实报实销"],
    ]
    for r, row in enumerate(rows, start=1):
        for c, v in enumerate(row):
            t.cell(r, c).text = v
    p = OUT / "费用报销管理办法.docx"
    d.save(str(p))
    return p


def make_xlsx() -> Path:
    """Excel：多 sheet。

    ★ 多 sheet 是常被漏掉的：一个"价格表.xlsx"可能有
      "华东/华北/华南"三个 sheet，只读第一个的话就丢了两大块。
    """
    import openpyxl

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "年假规则"
    ws.append(["工龄", "年假天数"])
    ws.append(["满1年", "5天"])
    ws.append(["满2年", "6天"])
    ws.append(["满3年", "7天"])

    ws2 = wb.create_sheet("请假流程")
    ws2.append(["步骤", "说明"])
    ws2.append(["1", "在企业微信提交申请"])
    ws2.append(["2", "直属主管审批"])
    ws2.append(["3", "超过3天需部门负责人审批"])

    p = OUT / "员工假期制度.xlsx"
    wb.save(str(p))
    return p


def make_pdf_text() -> Path:
    """有文本层的 PDF（正常情况）。"""
    from pypdf import PdfWriter
    from pypdf.generic import DecodedStreamObject, DictionaryObject, NameObject

    w = PdfWriter()
    page = w.add_blank_page(width=595, height=842)
    # 手工塞一段带文本层的流，模拟"电子版 PDF"
    font = DictionaryObject()
    font.update({
        NameObject("/Type"): NameObject("/Font"),
        NameObject("/Subtype"): NameObject("/Type1"),
        NameObject("/BaseFont"): NameObject("/Helvetica"),
    })
    res = DictionaryObject()
    fonts = DictionaryObject()
    fonts[NameObject("/F1")] = w._add_object(font)
    res[NameObject("/Font")] = fonts
    page[NameObject("/Resources")] = res

    content = (
        b"BT /F1 14 Tf 60 760 Td (Annual leave policy) Tj ET\n"
        b"BT /F1 12 Tf 60 730 Td (Employees get 5 days after one year.) Tj ET\n"
        b"BT /F1 12 Tf 60 700 Td (Each extra year adds 1 day, max 15 days.) Tj ET\n"
    )
    stream = DecodedStreamObject()
    stream.set_data(content)
    page[NameObject("/Contents")] = w._add_object(stream)
    p = OUT / "annual_leave.pdf"
    with open(p, "wb") as f:
        w.write(f)
    return p


def make_pdf_scanned() -> Path:
    """没有文本层的 PDF（扫描件）。

    ★ 这是**故意造的**：它模拟"扫描件/图片型 PDF"。
      正确行为是**如实说没有文本层**，而不是 OCR 硬猜。
      硬猜出来的字是错的，而**错的知识比没有知识更危险** ——
      AI 会拿着错字去回答，用户还以为是官方的。
    """
    from pypdf import PdfWriter

    w = PdfWriter()
    w.add_blank_page(width=595, height=842)  # 空白页 = 没有文本层
    p = OUT / "扫描件-合同扫描.pdf"
    with open(p, "wb") as f:
        w.write(f)
    return p


def make_txt_with_gbk() -> Path:
    """GBK 编码的 txt。

    ★ 中小企业里大量存在 GBK/GB18030 的文件（尤其是 Windows 老系统导出的）。
      只按 utf-8 读就会抛异常或者读出乱码 —— 而乱码进了索引，
      AI 会拿它去回答，用户看到一堆问号。
    """
    p = OUT / "内部通知.txt"
    p.write_bytes(
        "标题: 内部通知\n权限: external\n---\n本周五下午三点开季度会，请提前十分钟到会议室。\n"
        .encode("gbk")
    )
    return p


made = []
for fn in (make_docx, make_xlsx, make_pdf_text, make_pdf_scanned, make_txt_with_gbk):
    try:
        made.append(fn())
    except Exception as exc:  # noqa: BLE001
        print(f"  ❌ {fn.__name__}: {type(exc).__name__}: {exc}")

print(f"造了 {len(made)} 个文件到 {OUT}:")
for p in made:
    print(f"  {p.name}  {p.stat().st_size} 字节")
