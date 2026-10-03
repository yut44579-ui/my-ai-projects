"""导入接口测试：AC1 / AC2 / AC4 / AC7(后端半边) / AC8 / AC9 + 预览信息完整性。"""

from __future__ import annotations

import hashlib
import io
import json
import time

import pytest
from fastapi.testclient import TestClient
from openpyxl import Workbook
from sqlalchemy import select
from sqlalchemy.orm import Session

from app.models.customer import Customer
from app.models.import_batch import ImportBatch
from app.services.errors import ImportErrorCode, ImportFailure
from app.services.tabular import MAX_FILE_BYTES, MAX_ROWS, parse_table
from tests.conftest import commit, upload

MARK_EMAIL = "@import.test"
MARK_PHONE = "19999"


def _csv(*rows: str, header: str = "客户姓名,联系电话,邮箱") -> bytes:
    return ("\n".join([header, *rows]) + "\n").encode("utf-8")


# ─────────────── AC1 ───────────────


def test_ac1_clean_fixture_imports_eight_rows(
    client: TestClient, db: Session, clean_imports: None, clean_csv: tuple[str, bytes]
) -> None:
    """AC1：导入 8 行 CSV → 8 行入库；每条 batch_id 指向该批次，source_type 与上传声明一致。"""
    filename, raw = clean_csv
    before_total = client.get("/api/customers").json()["total"]["value"]

    preview = upload(client, filename, raw)
    assert preview.status_code == 200, preview.text
    preview_body = preview.json()

    assert preview_body["file_kind"] == "csv"
    assert preview_body["encoding"] == "utf-8"
    assert preview_body["rows_total"]["value"] == 8
    assert preview_body["headers"] == ["客户姓名", "联系电话", "邮箱", "公司名称", "地区", "备注"]
    # 6 列全部自动命中，无未映射列
    assert [c["target"] for c in preview_body["columns"]] == [
        "name",
        "phone",
        "email",
        "company_name",
        "region",
        "note",
    ]
    assert all(c["confidence"] == "auto" for c in preview_body["columns"])
    assert preview_body["unmapped_columns"] == []
    # 前 3 行样本
    assert len(preview_body["sample_rows"]) == 3
    assert set(preview_body["sample_rows"][0]) == set(preview_body["headers"])
    assert preview_body["file_sha256"]

    resp = commit(client, filename, raw, preview_body["file_sha256"])
    assert resp.status_code == 200, resp.text
    body = resp.json()

    assert body["status"] == "SUCCESS"
    assert body["source_type"] == "TEST"
    assert body["rows_total"]["value"] == 8
    assert body["rows_created"]["value"] == 8
    assert body["rows_deduplicated"]["value"] == 0
    assert body["rows_skipped"]["value"] == 0
    assert body["skipped_reasons"] == []
    assert body["skipped_columns"] == []
    batch_id = body["batch_id"]
    assert body["evidence_ref"] == f"import_batch:{batch_id}"

    # 计数恒等式（业务数字一律 EvidenceValue）
    assert body["rows_total"]["value"] == (
        body["rows_created"]["value"]
        + body["rows_deduplicated"]["value"]
        + body["rows_skipped"]["value"]
    )

    # 批次详情：该批次导入的客户 id 列表 + 每条都指向该批次
    detail = client.get(f"/api/imports/{batch_id}").json()
    assert len(detail["customer_ids"]) == 8
    assert all(c["source_type"] == "TEST" for c in detail["customers"])
    assert all(c["evidence_ref"] == f"import_batch:{batch_id}" for c in detail["customers"])

    in_db = db.execute(select(Customer).where(Customer.batch_id == batch_id)).scalars().all()
    assert len(in_db) == 8
    assert all(c.source_type.value == "TEST" for c in in_db)

    # 客户列表：总数 = 导入前 + 8，字段齐全
    listing = client.get("/api/customers", params={"page_size": 100}).json()
    assert listing["total"]["value"] == before_total + 8
    assert set(listing["total"]) >= {"value", "state", "source_type", "evidence_ref"}
    names = {item["name"] for item in listing["items"]}
    assert {"张伟", "李娜", "王强", "赵敏", "孙丽", "周杰", "吴涛", "郑霞"} <= names


# ─────────────── AC2 ───────────────


def test_ac2_unmapped_note_column_never_leaks_into_business_field(
    client: TestClient, db: Session, clean_imports: None
) -> None:
    """AC2：把备注列写成"备注信息XX" → skipped_columns 含该列，且 note 确实为空（反向验证）。"""
    raw = _csv(
        f"甲,1999900201,ac2{MARK_EMAIL},这段文字绝对不许进 note 字段",
        header="客户姓名,联系电话,邮箱,备注信息XX",
    )
    preview_body = upload(client, "ac2.csv", raw).json()
    assert preview_body["unmapped_columns"] == ["备注信息XX"]
    assert [c["confidence"] for c in preview_body["columns"]] == ["auto", "auto", "auto", "unknown"]

    body = commit(client, "ac2.csv", raw, preview_body["file_sha256"]).json()

    assert body["skipped_columns"] == [{"column": "备注信息XX", "reason": "unmapped_column"}]
    assert body["rows_created"]["value"] == 1

    customer = db.execute(select(Customer).where(Customer.email == f"ac2{MARK_EMAIL}")).scalars().one()
    assert customer.note is None, "未映射列的内容绝不能被硬塞进 note"
    assert customer.company_name is None


# ─────────────── AC4 ───────────────


def test_ac4_empty_file_is_rejected(client: TestClient, clean_imports: None) -> None:
    """AC4：空文件 → 400 + {"error":"empty_file"}，零入库。"""
    resp = upload(client, "empty.csv", b"")
    assert resp.status_code == 400, resp.text
    assert resp.json()["error"] == "empty_file"

    # commit 侧：sha256 对得上（就是同一份空文件）才轮到文件级校验，仍是 empty_file 且零入库
    resp2 = commit(client, "empty.csv", b"", hashlib.sha256(b"").hexdigest())
    assert resp2.status_code == 400
    assert resp2.json()["error"] == "empty_file"


def test_ac4_unsupported_extension_is_rejected(client: TestClient, clean_imports: None) -> None:
    """AC4：上传 .txt → 400 + {"error":"unsupported_format"}。"""
    resp = upload(client, "notes.txt", "客户姓名,电话\n甲,13800138000\n".encode("utf-8"))
    assert resp.status_code == 400, resp.text
    assert resp.json()["error"] == "unsupported_format"


def test_file_level_failures_leave_no_batch_behind(
    client: TestClient, db: Session, clean_imports: None
) -> None:
    """文件级失败 = 零入库：不落 import_batches 行，也不落客户行。"""
    before_batches = db.execute(select(ImportBatch.id)).scalars().all()
    before_customers = db.execute(select(Customer.id)).scalars().all()

    assert upload(client, "empty.csv", b"").status_code == 400
    assert upload(client, "x.txt", b"abc").status_code == 400

    # 结束读事务，避免 MySQL REPEATABLE READ 拿到旧快照
    db.rollback()
    assert db.execute(select(ImportBatch.id)).scalars().all() == before_batches
    assert db.execute(select(Customer.id)).scalars().all() == before_customers


def test_limits_too_large_and_too_many_rows() -> None:
    """超限：单文件 >10MB → too_large；数据行 >50,000 → too_many_rows。"""
    with pytest.raises(ImportFailure) as big:
        parse_table(b"a" * (MAX_FILE_BYTES + 1), "big.csv")
    assert big.value.code is ImportErrorCode.TOO_LARGE

    header = "客户姓名,联系电话,邮箱\n"
    row = "甲,1999900001,x@import.test\n"
    overflow = (header + row * (MAX_ROWS + 1)).encode("utf-8")
    with pytest.raises(ImportFailure) as many:
        parse_table(overflow, "many.csv")
    assert many.value.code is ImportErrorCode.TOO_MANY_ROWS


# ─────────────── AC7（后端半边）：取消某列 → 进 skipped_columns ───────────────


def test_ac7_cancelled_column_goes_to_skipped_columns(
    client: TestClient, db: Session, clean_imports: None
) -> None:
    """AC7：用户在弹窗里取消某列（target=null）→ 该列进 skipped_columns，字段确实为空。"""
    raw = _csv(f"甲,1999900301,ac7{MARK_EMAIL},这段备注被用户取消映射", header="客户姓名,联系电话,邮箱,备注")
    preview_body = upload(client, "ac7.csv", raw).json()
    assert [c["target"] for c in preview_body["columns"]][-1] == "note"  # 自动预填到了 note

    mapping = json.dumps(
        [
            {"column": "客户姓名", "target": "name"},
            {"column": "联系电话", "target": "phone"},
            {"column": "邮箱", "target": "email"},
            {"column": "备注", "target": None},  # ★ 用户点"忽略此列"
        ]
    )
    body = commit(client, "ac7.csv", raw, preview_body["file_sha256"], mapping=mapping).json()

    assert body["skipped_columns"] == [{"column": "备注", "reason": "unmapped_column"}]
    customer = db.execute(select(Customer).where(Customer.email == f"ac7{MARK_EMAIL}")).scalars().one()
    assert customer.note is None


def test_ambiguous_alias_conflict_is_rejected(client: TestClient, clean_imports: None) -> None:
    """别名冲突：两个列命中同一字段 → 明确报错，不许任选一个。"""
    raw = _csv("甲,1999900302,13800138000", header="客户姓名,电话,联系电话")
    resp = upload(client, "conflict_alias.csv", raw)
    assert resp.status_code == 400, resp.text
    body = resp.json()
    assert body["error"] == "ambiguous_mapping"
    assert "phone" in body["conflicts"]


def test_invalid_mapping_target_is_rejected(client: TestClient, clean_imports: None) -> None:
    raw = _csv(f"甲,1999900303,map{MARK_EMAIL}")
    sha = upload(client, "bad_mapping.csv", raw).json()["file_sha256"]
    mapping = json.dumps([{"column": "客户姓名", "target": "not_a_field"}])
    resp = commit(client, "bad_mapping.csv", raw, sha, mapping=mapping)
    assert resp.status_code == 400
    assert resp.json()["error"] == "invalid_mapping"


# ─────────────── AC8 ───────────────


def test_ac8_same_sha256_is_rejected_with_batch_id(
    client: TestClient, db: Session, clean_imports: None
) -> None:
    """AC8：同一文件（sha256 相同）重复提交 → 被拒绝并提示「该文件已导入过，batch_id=N」。"""
    raw = _csv(f"甲,1999900401,ac8{MARK_EMAIL}")
    sha = upload(client, "ac8.csv", raw).json()["file_sha256"]

    first = commit(client, "ac8.csv", raw, sha)
    assert first.status_code == 200
    batch_id = first.json()["batch_id"]

    second = commit(client, "ac8.csv", raw, sha)
    assert second.status_code == 409, second.text
    body = second.json()
    assert body["error"] == "duplicate_import"
    assert f"该文件已导入过，batch_id={batch_id}" in body["message"]
    assert body["batch_id"] == batch_id

    # 重复提交不许产生第二个批次
    same_sha = db.execute(
        select(ImportBatch).where(ImportBatch.file_sha256 == sha)
    ).scalars().all()
    assert len(same_sha) == 1


def test_sha256_mismatch_between_preview_and_commit_is_rejected(
    client: TestClient, clean_imports: None
) -> None:
    """★ 防"预览与提交之间文件被换"：sha256 对不上必须拒绝。"""
    raw = _csv(f"甲,1999900402,swap{MARK_EMAIL}")
    sha = upload(client, "swap.csv", raw).json()["file_sha256"]

    other = _csv(f"乙,1999900403,swapped{MARK_EMAIL}")
    resp = commit(client, "swap.csv", other, sha)
    assert resp.status_code == 400, resp.text
    assert resp.json()["error"] == "sha256_mismatch"


# ─────────────── AC9 ───────────────


def test_ac9_500_rows_import_under_5_seconds(
    client: TestClient, db: Session, clean_imports: None
) -> None:
    """AC9：导入 500 行 CSV 在 5 秒内完成（本地 MySQL）。"""
    rows = [f"批量客户{i},{MARK_PHONE}{i:06d},bulk{i}{MARK_EMAIL}" for i in range(500)]
    raw = _csv(*rows)

    started = time.perf_counter()
    sha = upload(client, "bulk500.csv", raw).json()["file_sha256"]
    body = commit(client, "bulk500.csv", raw, sha).json()
    elapsed = time.perf_counter() - started

    assert body["rows_created"]["value"] == 500
    assert body["rows_skipped"]["value"] == 0
    assert body["status"] == "SUCCESS"
    assert elapsed < 5.0, f"500 行导入耗时 {elapsed:.2f}s，超过 5s 上限"


# ─────────────── §四 / §五 的其它硬点 ───────────────


def test_messy_fixture_partial_with_skip_reasons_and_unmapped_columns(
    client: TestClient, db: Session, clean_imports: None, messy_csv: tuple[str, bytes]
) -> None:
    """messy 文件（GB18030）：PARTIAL + 三类计数 + 跳过原因 + 未映射列。"""
    filename, raw = messy_csv

    preview_body = upload(client, filename, raw).json()
    assert preview_body["encoding"] == "gb18030", "非 UTF-8 中文文件必须被正确识别"
    assert preview_body["rows_total"]["value"] == 9
    assert preview_body["unmapped_columns"] == ["备注信息XX", "多余列A"]

    body = commit(client, filename, raw, preview_body["file_sha256"]).json()

    assert body["status"] == "PARTIAL"
    assert body["rows_total"]["value"] == 9
    assert body["rows_created"]["value"] == 4
    assert body["rows_deduplicated"]["value"] == 1
    assert body["rows_skipped"]["value"] == 4
    assert body["rows_total"]["value"] == (
        body["rows_created"]["value"] + body["rows_deduplicated"]["value"] + body["rows_skipped"]["value"]
    )

    reasons = {item["reason"] for item in body["skipped_reasons"]}
    assert reasons == {"missing_name", "invalid_phone", "invalid_email", "empty_row"}
    raw_values = {item["reason"]: item["raw"] for item in body["skipped_reasons"]}
    assert raw_values["invalid_phone"] == "abc-def"
    assert raw_values["invalid_email"] == "not-an-email"
    # 行号必须是文件里的真实行号（含表头行），用户能直接回 Excel 定位
    assert {item["row"] for item in body["skipped_reasons"]} == {6, 7, 8, 10}

    assert [item["column"] for item in body["skipped_columns"]] == ["备注信息XX", "多余列A"]

    # GB18030 解码正确：中文没有乱码
    names = {c.name for c in db.execute(select(Customer).where(Customer.batch_id == body["batch_id"])).scalars()}
    assert names == {"张伟", "李娜", "王强", "周杰"}
    # 未映射列的内容没有混进 note
    assert all(c.note is None for c in db.execute(select(Customer).where(Customer.batch_id == body["batch_id"])).scalars())


def test_header_row_can_be_specified(client: TestClient, db: Session, clean_imports: None) -> None:
    """§四：表头不一定是第一行 —— 预览里指定第几行是表头。"""
    raw = (
        "导出时间：2026-10-03\n"
        "本文件由测试系统生成\n"
        f"客户姓名,联系电话,邮箱\n"
        f"甲,1999900501,header{MARK_EMAIL}\n"
    ).encode("utf-8")

    # 默认第 1 行当表头时结构对不上 → 明确报 parse_error，不许静默丢数据
    default = upload(client, "header.csv", raw)
    assert default.status_code == 400
    assert default.json()["error"] == "parse_error"

    moved = upload(client, "header.csv", raw, header_row=3).json()
    assert moved["header_row"] == 3
    assert moved["ignored_leading_rows"] == 2
    assert moved["headers"] == ["客户姓名", "联系电话", "邮箱"]
    assert moved["rows_total"]["value"] == 1

    body = commit(client, "header.csv", raw, moved["file_sha256"], header_row=3).json()
    assert body["rows_created"]["value"] == 1
    customer = db.execute(select(Customer).where(Customer.email == f"header{MARK_EMAIL}")).scalars().one()
    assert customer.phone == "1999900501"


def test_xlsx_multi_sheet_is_explicitly_reported(
    client: TestClient, db: Session, clean_imports: None
) -> None:
    """§四：xlsx 多 sheet 只读第一个时必须在响应里明确告知，不许静默。"""
    workbook = Workbook()
    first = workbook.active
    first.title = "客户明细"
    first.append(["客户姓名", "联系电话", "邮箱"])
    first.append(["甲", 1999900601, f"sheet1{MARK_EMAIL}"])  # ★ 电话是数字单元格
    second = workbook.create_sheet("废弃表")
    second.append(["不应该被读到"])
    buffer = io.BytesIO()
    workbook.save(buffer)
    workbook.close()
    raw = buffer.getvalue()

    preview_body = upload(client, "two_sheets.xlsx", raw).json()

    assert preview_body["file_kind"] == "xlsx"
    assert preview_body["sheet_names"] == ["客户明细", "废弃表"]
    assert preview_body["sheet_used"] == "客户明细"
    assert preview_body["sheet_note"] and "只读取" in preview_body["sheet_note"]
    assert preview_body["rows_total"]["value"] == 1

    body = commit(client, "two_sheets.xlsx", raw, preview_body["file_sha256"]).json()
    assert body["rows_created"]["value"] == 1
    customer = db.execute(select(Customer).where(Customer.email == f"sheet1{MARK_EMAIL}")).scalars().one()
    # ★ 强制 dtype=str：数字单元格不能被存成 1999900601.0 或科学计数
    assert customer.phone == "1999900601"


def test_xlsx_is_accepted_and_second_sheet_can_be_chosen(
    client: TestClient, clean_imports: None
) -> None:
    workbook = Workbook()
    workbook.active.title = "第一张"
    workbook.active.append(["客户姓名"])
    workbook.active.append(["甲"])
    other = workbook.create_sheet("第二张")
    other.append(["客户姓名", "联系电话"])
    other.append(["乙", "1999900602"])
    buffer = io.BytesIO()
    workbook.save(buffer)
    raw = buffer.getvalue()

    picked = upload(client, "pick.xlsx", raw, sheet="第二张").json()
    assert picked["sheet_used"] == "第二张"
    assert picked["headers"] == ["客户姓名", "联系电话"]

    missing = upload(client, "pick.xlsx", raw, sheet="不存在")
    assert missing.status_code == 400
    assert missing.json()["error"] == "parse_error"


def test_batch_list_and_detail_expose_counts_as_evidence(
    client: TestClient, db: Session, clean_imports: None
) -> None:
    """GET /api/imports 与 /api/imports/{id}：计数是业务数字 → 走 EvidenceValue。"""
    raw = _csv(f"甲,1999900701,list{MARK_EMAIL}")
    body = commit(client, "list.csv", raw, upload(client, "list.csv", raw).json()["file_sha256"]).json()
    batch_id = body["batch_id"]

    listing = client.get("/api/imports").json()
    assert listing["total"]["value"] >= 1
    item = next(i for i in listing["items"] if i["batch_id"] == batch_id)
    assert item["rows_created"]["value"] == 1
    assert set(item["rows_created"]) >= {"value", "state", "source_type", "evidence_ref"}

    detail = client.get(f"/api/imports/{batch_id}").json()
    assert detail["batch_id"] == batch_id
    assert detail["customer_ids"] == [c["id"] for c in detail["customers"]]
    assert client.get("/api/imports/99999999").status_code == 404
