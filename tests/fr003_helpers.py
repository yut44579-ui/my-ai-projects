"""fr003_helpers.py · FR-003 各测试文件共用的样例文件与隔离夹具。

════════════════════════════════════════════════════════════════════════
【为什么要有这个模块】
════════════════════════════════════════════════════════════════════════
FR-003 的六个测试文件都要造"一份真的 Excel / CSV / Markdown / PPT / Word / PDF"。
如果每个文件各写一遍，改一处忘一处，最后"测的到底是不是同一个东西"就说不清了。
这里集中一份，**而且全部造真文件**（真 xlsx、真 pptx、真 docx、真 pdf）：

    · 不 mock 解析器（mock 了就等于没测解析）
    · 不拿假字节糊弄（假文件会被签名/结构检查挡在门外，测出来的"通过"毫无意义）

Word / PDF 的造法**复用既有测试里已经验证过的那两个函数**
（`tests/test_documents.py::make_docx` / `build_minimal_pdf`）——
不重复造轮子，也保证"能过文档管道的文件"在这里是同一份。

════════════════════════════════════════════════════════════════════════
【隔离：每个用例都把六个路径指到 tmp_path】
════════════════════════════════════════════════════════════════════════
`isolate(tmp_path)` 把 SRA_STATE_DIR / SRA_DB_PATH / SRA_ORIGINAL_DIR /
SRA_UPLOAD_DIR / SRA_DOC_DIR / SRA_OUTPUT_DIR 全指向临时目录，
`reset()` 再幂等建表 + 清掉进程内的帧缓存。
两个一起调，用例之间**互不污染**，也永远碰不到真实的 `state/` 与 `data/app.db`。
"""

from __future__ import annotations

import datetime as _dt
import hashlib
import os
import pathlib
from typing import Any

from app.importer import db, frames

#: 一个用例要隔离的全部环境变量（少一个就可能漏进真实目录）
ISOLATED_ENV: tuple[str, ...] = (
    "SRA_STATE_DIR",
    "SRA_DB_PATH",
    "SRA_ORIGINAL_DIR",
    "SRA_UPLOAD_DIR",
    "SRA_DOC_DIR",
    "SRA_OUTPUT_DIR",
)


def isolate(root: pathlib.Path, monkeypatch: Any = None) -> dict[str, str]:
    """把六个路径指到 `root` 下面（每个用例一个临时目录）。返回实际用到的路径。

    传了 `monkeypatch` 就用它设（用例结束后**自动还原**）；没传就直接改 `os.environ`
    （脚本里用）。测试一律传 monkeypatch —— 谁也不想自己的临时目录漏给下一个用例。
    """
    paths = {
        "SRA_STATE_DIR": root / "state",
        "SRA_DB_PATH": root / "app.db",
        "SRA_ORIGINAL_DIR": root / "original",
        "SRA_UPLOAD_DIR": root / "uploads",
        "SRA_DOC_DIR": root / "documents",
        "SRA_OUTPUT_DIR": root / "outputs",
    }
    for key, path in paths.items():
        if monkeypatch is not None:
            monkeypatch.setenv(key, str(path))
        else:
            os.environ[key] = str(path)
    for key, path in paths.items():
        if key != "SRA_DB_PATH":
            path.mkdir(parents=True, exist_ok=True)
    return {key: str(path) for key, path in paths.items()}


def reset() -> None:
    """建表 + 清帧缓存（换了隔离目录就必须调一次）。"""
    db.init_schema()
    frames.clear_cache()


# ════════════════════════════════════════════════════════════════════════
# 样例文件（全部是真文件）
# ════════════════════════════════════════════════════════════════════════
def make_xlsx(path: pathlib.Path, sheets: dict[str, list[list[Any]]]) -> pathlib.Path:
    """造一个多工作表的 xlsx（第一行是表头）。"""
    import pandas as pd

    with pd.ExcelWriter(path, engine="openpyxl") as writer:
        for sheet, rows in sheets.items():
            frame = pd.DataFrame(rows[1:], columns=rows[0])
            frame.to_excel(writer, sheet_name=sheet, index=False)
    return path


def make_csv(path: pathlib.Path, text: str, encoding: str = "utf-8") -> pathlib.Path:
    path.write_text(text, encoding=encoding)
    return path


def make_md(path: pathlib.Path, text: str) -> pathlib.Path:
    path.write_text(text, encoding="utf-8")
    return path


def make_pptx(path: pathlib.Path, slides: list[dict[str, Any]]) -> pathlib.Path:
    """造一个 pptx。`slides` 里每项：`{"title": str, "body": str, "tables": [行列表]}`。

    表格用 python-pptx 真的建出来（不是往 zip 里塞 XML）—— 这样解析测的是真实结构。
    """
    from pptx import Presentation
    from pptx.util import Inches

    presentation = Presentation()
    layout = presentation.slide_layouts[5]           # 5 = 只有标题的版式
    for spec in slides:
        slide = presentation.slides.add_slide(layout)
        if spec.get("title"):
            slide.shapes.title.text = str(spec["title"])
        if spec.get("body"):
            box = slide.shapes.add_textbox(Inches(0.5), Inches(2), Inches(6), Inches(1))
            box.text_frame.text = str(spec["body"])
        for table_rows in spec.get("tables") or []:
            rows, columns = len(table_rows), len(table_rows[0])
            shape = slide.shapes.add_table(
                rows, columns, Inches(0.5), Inches(3), Inches(6), Inches(0.8 * rows)
            )
            for row_index, row in enumerate(table_rows):
                for column_index, value in enumerate(row):
                    shape.table.cell(row_index, column_index).text = str(value)
    presentation.save(str(path))
    return path


def make_docx(path: pathlib.Path, *, with_table: bool = True) -> pathlib.Path:
    """造一个真的 .docx（复用既有测试里那份实现 —— 不重复造轮子）。"""
    from test_documents import make_docx as _make

    _make(path, with_table=with_table)
    return path


def make_pdf(path: pathlib.Path, pages: list[str] | None = None) -> pathlib.Path:
    """造一个真的 .pdf（同样复用既有测试里的实现）。

    ★ 正文只能用 ASCII：那份最小 PDF 的内容流按 Latin-1 写（PDF 标准字体的默认编码），
      塞中文进去会在**造文件**这一步就报 UnicodeEncodeError —— 与解析能力无关。
    """
    from test_documents import build_minimal_pdf

    path.write_bytes(build_minimal_pdf(pages or ["Quarterly sales report"]))
    return path


# ════════════════════════════════════════════════════════════════════════
# 小工具
# ════════════════════════════════════════════════════════════════════════
def sha256_of(path: pathlib.Path) -> str:
    digest = hashlib.sha256()
    with open(path, "rb") as handle:
        for chunk in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def counts(source_file_id: str) -> dict[str, int]:
    """这份来源文件落成了几个数据集 / 几个文档（直接读库，不信接口的话）。"""
    from app.importer import store

    with db.readonly() as connection:
        return store.successful_import_count(connection, source_file_id)


def table_rows(dataset_id: str, table_name: str | None = None) -> list[dict[str, Any]]:
    """把物化的行读回来（核对"库里到底存了什么"用）。"""
    from app.importer import store

    with db.readonly() as connection:
        tables = store.list_dataset_tables(connection, dataset_id)
        chosen = table_name or (str(tables[0]["table_name"]) if tables else "")
        return store.load_rows(connection, dataset_id, chosen)


def all_imports() -> list[dict[str, Any]]:
    from app.importer import store

    with db.readonly() as connection:
        records, _total = store.list_imports(connection, limit=200)
    return records


#: 地区验收用的三行数据（评审 §五 D 段给的样例，一字不改）
REGION_SAMPLE = [["省份", "销售额"], ["广东", 100], ["浙江", 200], ["江苏", 300]]

#: 一份"正文 + 表格"的 Markdown
MD_SAMPLE = (
    "# 季度分析\n\n"
    "整体销售向好，下面的表格是分地区的金额。\n\n"
    "| 地区 | 销售额 |\n"
    "|---|---:|\n"
    "| 华东 | 100 |\n"
    "| 华南 | 200 |\n\n"
    "以上就是本季度的结论。\n"
)

#: 一份"正文 + 销售表 + 议程表"的 PPT（议程表**不该**被当成数据集）
PPTX_SAMPLE = [
    {
        "title": "上季度销售分析",
        "body": "整体向好，华东贡献最大。",
        "tables": [[["省份", "销售额"], ["广东", "100"], ["浙江", "200"]]],
    },
    {
        "title": "会议议程",
        "body": "",
        "tables": [[["时间", "事项"], ["10:00", "开场"], ["10:30", "讨论"]]],
    },
]


def today() -> _dt.date:
    return _dt.date(2011, 12, 9)
