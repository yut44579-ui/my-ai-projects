"""renderer.py · 渲染层：把**算好的数字**写进 Excel 报表，样式保持和模板一致（D5）。

职责边界：
    ✅ 用代码生成一份基础模板（openpyxl Workbook，**不经过 pandas**）
    ✅ 打开模板副本 → 只改 Spec 指定的数据单元格 → 另存到 outputs/
    ✅ 存完立刻读回来自检（值 + 数字格式），不对就删掉坏文件并报错（D10）
    ❌ 不取数、不算数（那是 loader.py / executor.py 的活，数字一律由代码算 —— D4）
    ❌ 不重建文件（`pandas.to_excel` 会丢合并单元格/字体/列宽/条件格式 —— D5 明确否掉）

为什么"原地改"这么重要（D5）：
    用户给参考物时的诉求是"报表要像这个样子"。pandas 重新生成的文件虽然数字对，
    但样式全丢（列宽、字体、数字格式、合并标题），对用户来说就是"不能用"。

Spec 与算数的衔接（两个桥接函数，别小看它们 —— 这是"口径固化"落地的地方）：
    executor_kwargs(spec)      Spec 里固化的排除规则（D16-3/4/5）→ compute_sales_amount(...) 的开关参数
    values_from_result(spec, result)
                               executor 的返回字典 → {指标名: 数值}
    这两个函数只做**白名单映射**（Spec 的受限 DSL ↔ executor 已产出的指标），
    遇到没定义的组合一律报错，绝不猜（CLAUDE.md 铁律 4）。
"""

from __future__ import annotations

import datetime as _dt
import time
from pathlib import Path
from typing import Any, Mapping

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils.cell import coordinate_to_tuple

from app.engine.metrics import EXCLUSION_RULES
from app.spec.models import ReportSpec

# ── 路径（本文件位于 <项目根>/app/engine/renderer.py，向上 3 层即项目根）────
PROJECT_ROOT = Path(__file__).resolve().parents[2]
OUTPUT_DIR = PROJECT_ROOT / "outputs"
DEFAULT_TEMPLATE_PATH = PROJECT_ROOT / "templates" / "weekly_sales_template.xlsx"

# ════════════════════════════════════════════════════════════════════════
# 基础模板的行定义（模板长什么样 + 默认单元格映射，一次定义两处共用）
# ════════════════════════════════════════════════════════════════════════
# 每行：(指标名, 报表标签, 数字格式, 口径说明)
# 指标名与 executor.compute_sales_amount(...) 的返回键对应（见 values_from_result 的白名单）
DEFAULT_METRIC_ROWS: tuple[tuple[str, str, str, str], ...] = (
    ("sales_amount", "销售额（元）", "#,##0.00", "sum(Quantity*UnitPrice)；按 D16 口径排除后求和"),
    ("rows_in_range", "区间原始行数", "#,##0", "含首尾全天的时间筛选结果，未做任何排除"),
    ("rows_valid", "有效行数", "#,##0", "区间原始行数 − 被排除行数"),
    ("rows_excluded", "被排除行数", "#,##0", "命中 ≥1 条已启用排除规则的行"),
    ("excluded_amount", "被排除金额（元）", "#,##0.00", "被排除行的 sum(Quantity*UnitPrice)（通常为负）"),
    ("valid_qty_sum", "有效数量合计", "#,##0", "有效行的 Quantity 之和"),
)

DEFAULT_SHEET_NAME = "周销售报表"
DEFAULT_FIRST_DATA_ROW = 4          # 第 1~2 行标题/副标题，第 3 行表头，数据从第 4 行开始

# 默认"指标名 → 单元格地址"（数据格固定在第 2 列 B，行号随上表顺序）
DEFAULT_CELL_MAP: dict[str, str] = {
    name: f"B{DEFAULT_FIRST_DATA_ROW + offset}"
    for offset, (name, *_) in enumerate(DEFAULT_METRIC_ROWS)
}

# 模板里的固定文案
_TITLE_TEXT = "周销售报表"
_SUBTITLE_TEXT = "口径：D16（含首尾全天；排除取消单 / 退货负数量 / 负零单价）｜数字一律由代码计算（D4）"
_HEADERS = ("指标", "数值", "口径说明")
_NOTE_TEXT = "本表由 sales-report-agent 自动生成：在模板上原地写入数据格，样式保持不变（D5）。"

# 样式（模板的"参考物风格"，渲染时不会动这些格）
_TITLE_FONT = Font(name="微软雅黑", size=14, bold=True, color="FF1F3864")
_SUBTITLE_FONT = Font(name="微软雅黑", size=9, italic=True, color="FF595959")
_HEADER_FONT = Font(name="微软雅黑", size=11, bold=True, color="FFFFFFFF")
_BODY_FONT = Font(name="微软雅黑", size=11)
_LABEL_FONT = Font(name="微软雅黑", size=11, bold=False)
_NOTE_FONT = Font(name="微软雅黑", size=9, italic=True, color="FF808080")

_TITLE_FILL = PatternFill("solid", fgColor="FFD9E1F2")
_HEADER_FILL = PatternFill("solid", fgColor="FF4472C4")
_ZEBRA_FILL = PatternFill("solid", fgColor="FFF2F6FC")

_THIN = Side(style="thin", color="FFBFBFBF")
_BORDER = Border(left=_THIN, right=_THIN, top=_THIN, bottom=_THIN)

_COLUMN_WIDTHS = {"A": 26.0, "B": 18.0, "C": 62.0}

# ════════════════════════════════════════════════════════════════════════
# 读回容差：openpyxl 把数字序列化成 "%.16g"（见 openpyxl/compat/strings.py 的 safe_string，
# 由 openpyxl/cell/_writer.py 调用），而双精度需要 **17 位**有效数字才能无损往返，
# 因此第 17 位可能差 1 个 ULP（相对误差 ≤ 2.3e-16）。这不是计算错误，是序列化损耗。
# 这里用 1e-9 的相对/绝对容差判"读回值 == 写入值"：
#   远大于 1 ULP（2.3e-16）→ 不会误报；远小于 1 分钱（1e-2）→ 抓得住真实的写错/写丢。
# ════════════════════════════════════════════════════════════════════════
_READBACK_REL_TOL = 1e-9
_READBACK_ABS_TOL = 1e-9

# ════════════════════════════════════════════════════════════════════════
# 受限 DSL ↔ executor 返回键 的白名单（没列出来的组合 = 当前引擎产不出的口径，报错不猜）
# ════════════════════════════════════════════════════════════════════════
# (op, field) → executor 返回字典里的取值路径（点号表示嵌套）
_SUPPORTED_METRICS: dict[tuple[str, str], str] = {
    ("sum", "sales_amount"): "amount",
    ("sum", "excluded_amount"): "excluded_amount",
    ("sum", "valid_qty_sum"): "valid_qty_sum",
    ("count", "rows_in_range"): "rows_in_range",
    ("count", "rows_valid"): "rows_valid",
    ("count", "rows_excluded"): "rows_excluded",
    ("count", "customer_id_nulls"): "validations.customer_id_nulls",
    ("count", "duplicate_rows"): "validations.duplicate_rows",
}


class RenderError(RuntimeError):
    """渲染失败（模板不存在 / sheet 不存在 / 单元格不可写 / 读回自检不过）。

    抛出即代表**不产出文件**（D10：失败不产出错误文档），调用方记错误、别交付。
    """


class UnsupportedMetricError(RenderError):
    """Spec 里的指标口径，当前引擎产不出来（受限 DSL 合法，但 executor 没有对应输出）。"""


# ════════════════════════════════════════════════════════════════════════
# 1. 生成基础模板（**用代码画，不用 pandas.to_excel** —— D5）
# ════════════════════════════════════════════════════════════════════════
def build_default_template(path: str | Path = DEFAULT_TEMPLATE_PATH, *, overwrite: bool = False) -> Path:
    """用 openpyxl 生成一份基础模板（标题 / 表头 / 指标行 / 样式），返回写出的路径。

    为什么不用 pandas.to_excel：D5 明确否掉 —— pandas 写出来的文件样式全丢，
    而"保留参考物样式"正是本项目要解决的问题本身。这里每一个字体/格式/列宽都是显式设置的。

    模板里**数据格是空的**（B4:B9），由 render_report 填；标签/表头/样式是"参考物"部分，渲染时不动。
    """
    target = Path(path)
    if target.exists() and not overwrite:
        raise RenderError(
            f"模板已存在，未覆盖：{target}（如需重建，显式传 overwrite=True）"
        )
    target.parent.mkdir(parents=True, exist_ok=True)

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = DEFAULT_SHEET_NAME

    # ── 第 1 行：标题（合并 + 居中 + 底色）──────────────────────────────
    sheet.cell(row=1, column=1, value=_TITLE_TEXT)
    sheet.merge_cells(start_row=1, start_column=1, end_row=1, end_column=3)
    sheet["A1"].font = _TITLE_FONT
    sheet["A1"].fill = _TITLE_FILL
    sheet["A1"].alignment = Alignment(horizontal="center", vertical="center")
    sheet.row_dimensions[1].height = 30

    # ── 第 2 行：副标题（口径说明，随模板一起固化）──────────────────────
    sheet.cell(row=2, column=1, value=_SUBTITLE_TEXT)
    sheet.merge_cells(start_row=2, start_column=1, end_row=2, end_column=3)
    sheet["A2"].font = _SUBTITLE_FONT
    sheet["A2"].alignment = Alignment(horizontal="left", vertical="center")
    sheet.row_dimensions[2].height = 18

    # ── 第 3 行：表头 ───────────────────────────────────────────────────
    for column, text in enumerate(_HEADERS, start=1):
        cell = sheet.cell(row=3, column=column, value=text)
        cell.font = _HEADER_FONT
        cell.fill = _HEADER_FILL
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = _BORDER
    sheet.row_dimensions[3].height = 22

    # ── 第 4 行起：指标行（A 列标签、B 列数据格留空、C 列口径说明）──────
    for offset, (_, label, number_format, note) in enumerate(DEFAULT_METRIC_ROWS):
        row = DEFAULT_FIRST_DATA_ROW + offset
        label_cell = sheet.cell(row=row, column=1, value=label)
        label_cell.font = _LABEL_FONT
        label_cell.alignment = Alignment(horizontal="left", vertical="center")
        label_cell.border = _BORDER
        if offset % 2 == 1:                       # 隔行浅色底，读起来不串行
            label_cell.fill = _ZEBRA_FILL

        data_cell = sheet.cell(row=row, column=2)  # ← 数据格：留空，等渲染填
        data_cell.number_format = number_format
        data_cell.font = _BODY_FONT
        data_cell.alignment = Alignment(horizontal="right", vertical="center")
        data_cell.border = _BORDER
        if offset % 2 == 1:
            data_cell.fill = _ZEBRA_FILL

        note_cell = sheet.cell(row=row, column=3, value=note)
        note_cell.font = Font(name="微软雅黑", size=9, color="FF595959")
        note_cell.alignment = Alignment(horizontal="left", vertical="center", wrap_text=False)
        note_cell.border = _BORDER
        if offset % 2 == 1:
            note_cell.fill = _ZEBRA_FILL

    # ── 末尾说明行 ──────────────────────────────────────────────────────
    note_row = DEFAULT_FIRST_DATA_ROW + len(DEFAULT_METRIC_ROWS) + 1
    sheet.cell(row=note_row, column=1, value=_NOTE_TEXT)
    sheet.merge_cells(start_row=note_row, start_column=1, end_row=note_row, end_column=3)
    sheet[f"A{note_row}"].font = _NOTE_FONT
    sheet[f"A{note_row}"].alignment = Alignment(horizontal="left", vertical="center")

    # ── 列宽 + 冻结表头 ────────────────────────────────────────────────
    for column, width in _COLUMN_WIDTHS.items():
        sheet.column_dimensions[column].width = width
    sheet.freeze_panes = f"A{DEFAULT_FIRST_DATA_ROW}"

    workbook.save(target)
    return target


# ════════════════════════════════════════════════════════════════════════
# 2. Spec ↔ executor 的桥接（口径固化落地处）
# ════════════════════════════════════════════════════════════════════════
def executor_kwargs(spec: ReportSpec) -> dict[str, bool]:
    """Spec 固化的排除规则 → compute_sales_amount(...) 的开关参数。

    参数名从 metrics.EXCLUSION_RULES[].param 取（口径单一来源，不在这里手写字符串），
    所以规则表改了这里自动跟着改，不会出现"两处口径漂移"。
    """
    enabled = set(spec.exclusion_keys())
    return {rule.param: rule.key in enabled for rule in EXCLUSION_RULES}


def values_from_result(spec: ReportSpec, result: Mapping[str, Any]) -> dict[str, int | float]:
    """把 executor 的返回字典，按 Spec 的受限 DSL 抽成 {指标名: 数值}。

    只认白名单里的 (op, field) 组合；遇到产不出的口径直接报错（附上可用组合），不猜。
    """
    values: dict[str, int | float] = {}
    for metric in spec.metrics:
        field = metric.dsl.field or ""
        key = (metric.dsl.op, field)
        if key not in _SUPPORTED_METRICS:
            supported = ", ".join(f"{op}({fld})" for op, fld in sorted(_SUPPORTED_METRICS))
            raise UnsupportedMetricError(
                f"指标 {metric.name!r} 的口径 op={metric.dsl.op!r} field={field!r} 当前引擎产不出来。\n"
                f"  当前支持：{supported}\n"
                f"  说明：口径必须先定义再计算（D12）—— 需要在 executor 里先落地该口径，才能写进 Spec。"
            )
        value = _lookup(result, _SUPPORTED_METRICS[key])
        if isinstance(value, bool) or not isinstance(value, (int, float)):
            raise RenderError(
                f"指标 {metric.name!r} 取到的值不是数字：{value!r}"
                f"（路径 {_SUPPORTED_METRICS[key]}；报表单元格只接受数字）"
            )
        values[metric.name] = value
    return values


def _lookup(result: Mapping[str, Any], path: str) -> Any:
    """按点号路径从 executor 的返回字典里取值（如 'validations.customer_id_nulls'）。"""
    node: Any = result
    for part in path.split("."):
        if not isinstance(node, Mapping) or part not in node:
            raise RenderError(f"executor 返回结果里找不到路径 {path!r}（缺 {part!r}）")
        node = node[part]
    return node


# ════════════════════════════════════════════════════════════════════════
# 3. 渲染：打开模板 → 只改数据格 → 另存 outputs/ → 读回自检
# ════════════════════════════════════════════════════════════════════════
def render_report(
    spec: ReportSpec,
    values: Mapping[str, int | float],
    *,
    output_path: str | Path | None = None,
    create_template_if_missing: bool = False,
) -> dict:
    """把 {指标名: 数值} 写进模板副本，另存到 outputs/，返回产出信息。

    参数：
        spec                          报表规格（决定模板/sheet/哪些格写哪个指标）
        values                        已算好的数字：{指标名: 数值}（由 executor 产出，见 values_from_result）
        output_path                   指定输出路径（默认为 outputs/<spec_id>_v<版>_<时间戳>.xlsx）
        create_template_if_missing    模板缺失时是否用代码生成基础模板：
                                      · 默认 False —— 遵循 D10（缺模板是配置错误，报错而不是悄悄换个模板）
                                      · 传 True  —— 按 TASK-002 第 9 条的便利行为，用 build_default_template 生成

    返回：
        {
          "output_path": str, "size_bytes": int, "template": str, "sheet": str,
          "spec_id", "spec_version",
          "cells_written": {指标名: {"cell","value","read_back","abs_diff","number_format"}},
          "verification": {"passed", "items", "rel_tol", "abs_tol"},
          "seconds": float,
        }

    抛错：
        UnsupportedMetricError  指标口径当前引擎产不出（见 values_from_result 白名单）
        RenderError             模板/sheet/单元格有问题，或读回自检不过（此时**不留下坏文件**）
    """
    started = time.perf_counter()

    # ── 前置校验（都在动文件之前做完 —— 失败绝不产出半成品）─────────────
    if spec.group_by:
        raise RenderError(
            f"group_by={list(spec.group_by)} 的**明细表**渲染方式尚未定义：本 TASK 只支持"
            f"「指标 → 单元格」的汇总表（模板里预留了格子）。分组明细需要先定表格布局口径（D12）。"
        )

    template = _resolve_path(spec.output.template)
    if not template.exists():
        if create_template_if_missing:
            build_default_template(template)
        else:
            raise RenderError(
                f"模板不存在：{template}\n"
                f"  可选：① 修正 spec.output.template；② 调用 renderer.build_default_template() "
                f"生成基础模板；③ 传 create_template_if_missing=True。"
            )

    cell_map = spec.output.cells
    if not cell_map:
        raise RenderError("spec.output.cells 为空：没有任何要写的单元格，Spec 不完整")
    missing_values = sorted(name for name in cell_map if name not in values)
    if missing_values:
        raise RenderError(
            f"缺少这些指标的数值：{missing_values}（已提供：{sorted(values)}）—— 不猜、不填 0"
        )

    workbook = load_workbook(template)
    if spec.output.sheet not in workbook.sheetnames:
        raise RenderError(
            f"模板里没有 sheet {spec.output.sheet!r}：{template}\n"
            f"  可选 sheet：{workbook.sheetnames}"
        )
    sheet = workbook[spec.output.sheet]

    # 模板的已用范围在**写之前**取：openpyxl 取单元格对象时会往内部表里插行，之后再取就不准了
    template_max_row, template_max_column = sheet.max_row, sheet.max_column

    # ── 只改数据格（其余单元格连样式都不碰）─────────────────────────────
    written: dict[str, dict[str, Any]] = {}
    for metric_name, address in cell_map.items():
        value = _as_number(values[metric_name], metric_name)
        _ensure_cell_writable(sheet, address, template_max_row, template_max_column)
        cell = sheet[address]
        number_format_before = cell.number_format      # 模板里预设的数字格式（渲染后必须还在）
        cell.value = value
        written[metric_name] = {
            "cell": address,
            "value": value,
            "number_format": number_format_before,
        }

    out_path = _resolve_output_path(spec, output_path)
    out_path.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(out_path)

    # ── 读回自检（D10：宁可报错，也不交付一个"看起来对"的文件）──────────
    verification = _verify_output(out_path, spec.output.sheet, written)

    return {
        "output_path": str(out_path),
        "size_bytes": out_path.stat().st_size,
        "template": str(template),
        "sheet": spec.output.sheet,
        "spec_id": spec.spec_id,
        "spec_version": spec.version,
        "cells_written": written,
        "verification": verification,
        "seconds": time.perf_counter() - started,
    }


def _as_number(value: Any, metric_name: str) -> int | float:
    """报表数据格只接受真数字（bool 不算 —— True/False 写进报表是错误用法）。"""
    if isinstance(value, bool) or not isinstance(value, (int, float)):
        raise RenderError(f"指标 {metric_name!r} 的值不是数字：{value!r}（只接受 int/float）")
    return value


def _ensure_cell_writable(sheet, address: str, template_max_row: int, template_max_column: int) -> None:
    """写之前确认格子可写：① 不在合并区（非左上角）；② 落在模板已用范围内。

    ② 是防"格子地址打错"的：模板只给 B4:B9 留了数据格，写到 B40 说明 Spec 或模板对不上，
       这种错误静默写进去= 报表少一个数还看不出来，所以要当场报错。
    """
    row, column = coordinate_to_tuple(address)
    for merged in sheet.merged_cells.ranges:
        if address in merged and (row, column) != (merged.min_row, merged.min_col):
            raise RenderError(
                f"单元格 {address} 在合并区 {merged.coord} 内且不是左上角，写入会被 Excel 忽略"
                f"（请改写到合并区左上角，或换一个数据格）"
            )
    if row > template_max_row or column > template_max_column:
        raise RenderError(
            f"单元格 {address} 超出模板已用范围（{sheet.title}!A1:"
            f"{_column_letter(template_max_column)}{template_max_row}）："
            f"模板里没有为它预留数据格，请核对 spec.output.cells 与模板"
        )


def _verify_output(path: Path, sheet_name: str, written: Mapping[str, Mapping[str, Any]]) -> dict:
    """把刚存的文件**重新读回来**逐个核对（数值 + 数字格式）；不对就删文件并报错。

    这是 D10 在渲染层的落地：出了坏文件，宁可没有文件。
    """
    workbook = load_workbook(path)
    sheet = workbook[sheet_name]
    items: list[dict[str, Any]] = []
    for metric_name, info in written.items():
        cell = sheet[info["cell"]]
        read_back = cell.value
        number_format_after = cell.number_format
        expected = info["value"]

        if isinstance(read_back, (int, float)) and not isinstance(read_back, bool):
            abs_diff = abs(float(read_back) - float(expected))
            tolerance = max(_READBACK_ABS_TOL, abs(float(expected)) * _READBACK_REL_TOL)
            value_passed = abs_diff <= tolerance
        else:
            abs_diff = None
            tolerance = None
            value_passed = read_back == expected       # 理论上不该走到这（只写数字）

        style_passed = number_format_after == info["number_format"]
        items.append(
            {
                "metric": metric_name,
                "cell": info["cell"],
                "expected": expected,
                "read_back": read_back,
                "abs_diff": abs_diff,
                "tolerance": tolerance,
                "value_passed": value_passed,
                "number_format": info["number_format"],
                "number_format_after": number_format_after,
                "style_passed": style_passed,
                "passed": bool(value_passed and style_passed),
            }
        )

    failed = [item for item in items if not item["passed"]]
    if failed:
        try:                                  # 坏文件不留在 outputs/ 里（D10）
            path.unlink()
        except OSError:
            pass
        detail = "; ".join(
            f"{item['cell']}({item['metric']}): 写入={item['expected']!r} 读回={item['read_back']!r} "
            f"格式 {item['number_format']!r}→{item['number_format_after']!r}"
            for item in failed
        )
        raise RenderError(f"读回自检不通过，已删除产出文件 {path}：{detail}")

    return {
        "passed": True,
        "items": items,
        "rel_tol": _READBACK_REL_TOL,
        "abs_tol": _READBACK_ABS_TOL,
    }


# ════════════════════════════════════════════════════════════════════════
# 路径小工具
# ════════════════════════════════════════════════════════════════════════
def _resolve_path(raw: str | Path) -> Path:
    """相对路径按**项目根**解析（Spec 里存的是可移植的相对路径，如 templates/xxx.xlsx）。"""
    path = Path(raw)
    return path if path.is_absolute() else PROJECT_ROOT / path


def _resolve_output_path(spec: ReportSpec, output_path: str | Path | None) -> Path:
    """决定产出文件路径：默认 outputs/<spec_id>_v<版本>_<时间戳>.xlsx。"""
    if output_path is not None:
        return _resolve_path(output_path)
    stamp = _dt.datetime.now().strftime("%Y%m%d-%H%M%S")
    return OUTPUT_DIR / f"{spec.spec_id}_v{spec.version}_{stamp}.xlsx"


def _column_letter(index: int) -> str:
    """列序号 → 列字母（1→A, 27→AA），只用于错误信息里显示模板范围。"""
    letters = ""
    while index > 0:
        index, remainder = divmod(index - 1, 26)
        letters = chr(ord("A") + remainder) + letters
    return letters
