"""export.py · 把**同一个查询结果**写成 xlsx / csv（STEP A）。

════════════════════════════════════════════════════════════════════════
【同源：导出不允许另算一遍】
════════════════════════════════════════════════════════════════════════
本文件**只做排版**，取数一律调 `queries.fetch_table(..., export=True)`——
与页面分页用的**是同一个函数、同一组筛选条件、同一套口径**，区别只有一个：导出要全部行。
所以"页面显示 A、导出文件里是 B"这件事在本设计里不可能发生（除非有人另外写一条取数路径）。

【两种格式】
  · xlsx：真表格（加粗表头 / 冻结首行 / 自动筛选 / 列宽 / 数字与金额格式 / 日期格式）
  · csv ：UTF-8 **带 BOM**（Excel 双击打开不乱码）+ CRLF 行尾

【不写进文件的东西】
不上报哈希、文件编号、路径这类审计字段：导出的是**业务表**，用户拿去就能用。
"""

from __future__ import annotations

import csv
import datetime as _dt
import io
from typing import Any

from app.datasets import queries

# 格式白名单
EXPORT_FORMATS: tuple[str, ...] = ("xlsx", "csv")
EXPORT_MIME: dict[str, str] = {
    "xlsx": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    "csv": "text/csv; charset=utf-8",
}
EXPORT_SUFFIX: dict[str, str] = {"xlsx": "xlsx", "csv": "csv"}

# 表头样式（与项目深色规格无关：导出的是给 Excel 看的文件，用白底+浅灰表头）
_HEADER_FILL = "F2F4F7"
_MONEY_FORMAT = "#,##0.00"
_INT_FORMAT = "#,##0"
_PCT_FORMAT = "0.00"

_FORMAT_MAP = {
    "money": _MONEY_FORMAT,
    "int": _INT_FORMAT,
    "qty": _INT_FORMAT,
    "pct": _PCT_FORMAT,
    "date": "yyyy-mm-dd",
    "datetime": "yyyy-mm-dd hh:mm:ss",
    "text": "@",
}


def export_filename(table: str, dataset: dict[str, Any], query: dict[str, Any], fmt: str) -> str:
    """导出文件名：业务可读（表名 + 数据源 + 区间），不带哈希/编号。"""
    label = queries.TABLE_LABELS.get(table, table)
    parts = [f"{dataset.get('name', '数据')}-{label}"]
    start, end = query.get("start"), query.get("end")
    if start and end:
        parts.append(f"{start}_{end}")
    stamp = _dt.datetime.now().strftime("%Y%m%d%H%M")
    return f"{'-'.join(parts)}-{stamp}.{EXPORT_SUFFIX[fmt]}"


def _cell_text(value: Any) -> str:
    """CSV 单元格：None → 空串（不是 "None"），时间 → 原样 ISO。"""
    if value is None:
        return ""
    if isinstance(value, bool):
        return "是" if value else "否"
    return str(value)


def to_csv(payload: dict[str, Any]) -> bytes:
    """表头 + 行 → CSV 字节（UTF-8 BOM + CRLF，Excel 双击直接打开）。"""
    columns = payload["columns"]
    buffer = io.StringIO()
    writer = csv.writer(buffer, lineterminator="\r\n")
    writer.writerow([column["label"] for column in columns])
    for item in payload["items"]:
        writer.writerow([_cell_text(item.get(column["key"])) for column in columns])
    return buffer.getvalue().encode("utf-8-sig")


def to_xlsx(payload: dict[str, Any]) -> bytes:
    """表头 + 行 → xlsx 字节（带格式；表头冻结、可筛选、列宽按内容估）。"""
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    columns = payload["columns"]
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = queries.TABLE_LABELS.get(payload["table"], "数据")[:31] or "数据"

    header_font = Font(bold=True)
    header_fill = PatternFill("solid", fgColor=_HEADER_FILL)
    for index, column in enumerate(columns, start=1):
        cell = sheet.cell(row=1, column=index, value=column["label"])
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center", vertical="center")
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = f"A1:{get_column_letter(len(columns))}1"

    widths = [len(str(column["label"])) * 2 for column in columns]
    for row_index, item in enumerate(payload["items"], start=2):
        for column_index, column in enumerate(columns, start=1):
            value = item.get(column["key"])
            column_format = column.get("format", "text")
            cell = sheet.cell(row=row_index, column=column_index,
                              value=_excel_value(value, column_format))
            number_format = _FORMAT_MAP.get(column_format)
            if number_format:
                cell.number_format = number_format
            if column.get("align") == "right":
                cell.alignment = Alignment(horizontal="right")
            widths[column_index - 1] = max(widths[column_index - 1], min(40, len(_cell_text(value)) + 4))
    for index, width in enumerate(widths, start=1):
        sheet.column_dimensions[get_column_letter(index)].width = min(48, max(10, width))

    stream = io.BytesIO()
    workbook.save(stream)
    return stream.getvalue()


def _excel_value(value: Any, fmt: str = "text") -> Any:
    """Excel 单元格值：**列声明为 date/datetime** 时，把日期字符串转成真日期对象。

    ★ 这里必须**跟着列的格式走**，不能见着像日期的字符串就转（FR-002B 的教训）：
      上一版不看列的格式，只要字符串长得像 `YYYY-MM-DD` 就转成 datetime，而转出来的
      单元格又挂着该列的 number_format —— 「期间」列当时声明的是文本（`@`），于是
      Excel 拿到"日期序列号 + 文本格式"这对组合，单元格里显示的就是 `40875`
      而不是 `2011-11-28`（用户实测反馈的那个缺陷）。

      所以：`date` / `datetime` 列 → 转真日期（配 `yyyy-mm-dd` 格式，读写都是日期类型）；
      其余列（文本/数字）→ **原样**，绝不把文本列偷偷变成日期。
    """
    if fmt not in ("date", "datetime"):
        return value
    if isinstance(value, str) and len(value) >= 10 and value[4:5] == "-" and value[7:8] == "-":
        try:
            return _dt.datetime.fromisoformat(value)
        except ValueError:
            return value                       # 认不出来的原样写，不猜、不崩
    return value


def build_export(table: str, fmt: str, **params: Any) -> tuple[str, bytes, str]:
    """导出入口：取数（与页面同源）→ 排版 → `(文件名, 字节, mime)`。"""
    if fmt not in EXPORT_FORMATS:
        raise queries.QueryError("unknown_format", f"导出格式只支持 {list(EXPORT_FORMATS)}，收到 {fmt!r}")
    payload = queries.fetch_table(table, export=True, **params)
    filename = export_filename(table, payload["dataset"], payload["query"], fmt)
    content = to_xlsx(payload) if fmt == "xlsx" else to_csv(payload)
    return filename, content, EXPORT_MIME[fmt]


__all__ = [
    "EXPORT_FORMATS",
    "EXPORT_MIME",
    "build_export",
    "export_filename",
    "to_csv",
    "to_xlsx",
]
