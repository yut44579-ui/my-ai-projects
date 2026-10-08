"""CSV / XLSX 解析为「字符串矩阵」（TASK-001 §四、§五）。

硬边界：只做 CSV/XLSX。依赖只用标准库 csv + openpyxl（xlsx 是 zip+XML 二进制容器，
标准库没有解析器，用成熟库而不是自己造轮子）。

★ 一律 dtype=str：电话 13800138000 被 Excel 存成数字会变科学计数、以 0 开头的会丢前导零，
  所以所有单元格统一转成字符串，禁止任何类型推断（本实现根本不引 pandas）。
"""

from __future__ import annotations

import csv
import hashlib
import io
import zipfile
from dataclasses import dataclass, field
from datetime import date, datetime
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.utils.exceptions import InvalidFileException

from app.services.encoding import decode_text
from app.services.errors import ImportErrorCode, ImportFailure

MAX_FILE_BYTES = 10 * 1024 * 1024  # 单文件 ≤ 10MB
MAX_ROWS = 50_000                  # 数据行 ≤ 50,000

CSV_SUFFIXES = {".csv"}
XLSX_SUFFIXES = {".xlsx", ".xlsm"}


@dataclass
class TableData:
    """解析结果：一个纯字符串矩阵 + 解析过程元信息（要如实回传给用户确认）。"""

    headers: list[str]
    rows: list[list[str]]                  # 已按表头宽度补齐
    file_kind: str                         # "csv" | "xlsx"
    header_row: int                        # 1 起的表头行号
    encoding: str | None                   # csv: 实际编码；xlsx: None
    encoding_note: str | None = None
    sheet_names: list[str] = field(default_factory=list)
    sheet_used: str | None = None
    sheet_note: str | None = None
    ignored_leading_rows: int = 0          # 表头之前被跳过的行数


def sha256_of(raw: bytes) -> str:
    return hashlib.sha256(raw).hexdigest()


def _check_size(raw: bytes) -> None:
    if not raw:
        raise ImportFailure(ImportErrorCode.EMPTY_FILE, "文件是空的（0 字节）")
    if len(raw) > MAX_FILE_BYTES:
        raise ImportFailure(
            ImportErrorCode.TOO_LARGE,
            f"文件 {len(raw) / 1024 / 1024:.1f}MB 超过上限 10MB",
        )


def _cell_to_str(value: object) -> str:
    """单元格 → 字符串（严禁类型推断导致的科学计数/丢前导零）。"""
    if value is None:
        return ""
    if isinstance(value, str):
        return value.strip()
    if isinstance(value, bool):
        return "TRUE" if value else "FALSE"
    if isinstance(value, int):
        return str(value)
    if isinstance(value, float):
        # Excel 里 13800138000 读出来是 13800138000.0 —— 整数浮点还原成整数写法
        return str(int(value)) if value.is_integer() else repr(value)
    if isinstance(value, (datetime, date)):
        return value.isoformat(sep=" ")
    return str(value).strip()


def _build_table(
    matrix: list[list[str]],
    *,
    header_row: int,
    file_kind: str,
    encoding: str | None,
    encoding_note: str | None = None,
    sheet_names: list[str] | None = None,
    sheet_used: str | None = None,
    sheet_note: str | None = None,
) -> TableData:
    """把原始行矩阵切成「表头 + 数据行」，并做结构与规模校验。"""
    if header_row < 1:
        raise ImportFailure(ImportErrorCode.PARSE_ERROR, f"表头行号必须 ≥ 1，收到 {header_row}")
    if len(matrix) < header_row:
        raise ImportFailure(
            ImportErrorCode.PARSE_ERROR,
            f"指定第 {header_row} 行是表头，但文件只有 {len(matrix)} 行",
        )

    headers = [c.strip() for c in matrix[header_row - 1]]
    if not any(headers):
        raise ImportFailure(ImportErrorCode.PARSE_ERROR, f"第 {header_row} 行（表头行）整行为空")

    # 重名列无法可靠映射 —— 明确报错，不许任选一列
    seen: set[str] = set()
    dupes: list[str] = []
    for name in headers:
        if name and name in seen:
            dupes.append(name)
        seen.add(name)
    if dupes:
        raise ImportFailure(
            ImportErrorCode.PARSE_ERROR,
            f"表头存在重名列：{', '.join(sorted(set(dupes)))}，无法可靠映射",
        )

    width = len(headers)
    data_rows: list[list[str]] = []
    for offset, raw_row in enumerate(matrix[header_row:], start=1):
        if len(raw_row) > width:
            extra = [c for c in raw_row[width:] if c]
            if extra:
                raise ImportFailure(
                    ImportErrorCode.PARSE_ERROR,
                    f"第 {header_row + offset} 行比表头多出 {len(raw_row) - width} 列"
                    f"（内容：{'/'.join(extra[:3])}），文件结构不一致",
                )
            raw_row = raw_row[:width]
        elif len(raw_row) < width:
            raw_row = list(raw_row) + [""] * (width - len(raw_row))
        data_rows.append(list(raw_row))
        if len(data_rows) > MAX_ROWS:
            raise ImportFailure(
                ImportErrorCode.TOO_MANY_ROWS,
                f"数据行超过上限 {MAX_ROWS} 行，请拆分后再导入",
            )

    return TableData(
        headers=headers,
        rows=data_rows,
        file_kind=file_kind,
        header_row=header_row,
        encoding=encoding,
        encoding_note=encoding_note,
        sheet_names=sheet_names or [],
        sheet_used=sheet_used,
        sheet_note=sheet_note,
        ignored_leading_rows=header_row - 1,
    )


def _parse_csv(raw: bytes, header_row: int, encoding: str | None) -> TableData:
    text, used_encoding = decode_text(raw, encoding)
    if not text.strip():
        raise ImportFailure(ImportErrorCode.EMPTY_FILE, "CSV 文件没有任何内容")

    reader = csv.reader(io.StringIO(text, newline=""))
    matrix = [[cell.strip().lstrip("﻿") for cell in row] for row in reader]
    # 丢掉尾部/中间的物理空行（csv.reader 对空行返回 []），它们不是"数据行"
    matrix = [row for row in matrix if row]

    note = None
    if used_encoding in {"gb18030"}:
        note = "按 GB18030/GBK 解码（文件不是 UTF-8，中文已正确还原）"

    return _build_table(
        matrix,
        header_row=header_row,
        file_kind="csv",
        encoding=used_encoding,
        encoding_note=note,
    )


def _parse_xlsx(raw: bytes, header_row: int, sheet: str | None) -> TableData:
    if not zipfile.is_zipfile(io.BytesIO(raw)):
        # 老式 .xls / 加密工作簿都不是 zip 容器，openpyxl 也读不了
        raise ImportFailure(
            ImportErrorCode.UNSUPPORTED_FORMAT,
            "该文件不是 xlsx（可能是老式 .xls 或已加密的工作簿），请另存为 .xlsx 或 .csv",
        )

    try:
        workbook = load_workbook(io.BytesIO(raw), read_only=True, data_only=True)
    except (InvalidFileException, KeyError, OSError, ValueError) as exc:
        raise ImportFailure(
            ImportErrorCode.PARSE_ERROR, f"xlsx 解析失败：{type(exc).__name__}: {exc}"
        ) from exc

    try:
        sheet_names = list(workbook.sheetnames)
        if not sheet_names:
            raise ImportFailure(ImportErrorCode.PARSE_ERROR, "xlsx 里没有任何 sheet")

        if sheet is not None:
            if sheet not in sheet_names:
                raise ImportFailure(
                    ImportErrorCode.PARSE_ERROR,
                    f"指定的 sheet {sheet!r} 不存在，文件里有：{', '.join(sheet_names)}",
                )
            used = sheet
        else:
            used = sheet_names[0]

        worksheet = workbook[used]
        matrix = [
            [_cell_to_str(cell) for cell in row]
            for row in worksheet.iter_rows(values_only=True)
        ]
        matrix = [row for row in matrix if any(cell for cell in row)]
    finally:
        workbook.close()

    # ★ 多 sheet 时必须明确告知，不许静默只读第一个
    if len(sheet_names) > 1:
        sheet_note = (
            f"文件含 {len(sheet_names)} 个 sheet：{'、'.join(sheet_names)}；"
            f"本次只读取「{used}」，其余未导入"
        )
    else:
        sheet_note = f"文件只有 1 个 sheet：「{used}」"

    return _build_table(
        matrix,
        header_row=header_row,
        file_kind="xlsx",
        encoding=None,
        encoding_note="xlsx 是二进制容器，不存在文本编码问题（内部 XML 为 UTF-8）",
        sheet_names=sheet_names,
        sheet_used=used,
        sheet_note=sheet_note,
    )


def parse_table(
    raw: bytes,
    filename: str,
    *,
    header_row: int = 1,
    sheet: str | None = None,
    encoding: str | None = None,
) -> TableData:
    """解析上传文件（CSV/XLSX）为字符串矩阵。文件级失败抛 ImportFailure。"""
    _check_size(raw)

    suffix = Path(filename or "").suffix.lower()
    if suffix in CSV_SUFFIXES:
        return _parse_csv(raw, header_row, encoding)
    if suffix in XLSX_SUFFIXES:
        return _parse_xlsx(raw, header_row, sheet)

    raise ImportFailure(
        ImportErrorCode.UNSUPPORTED_FORMAT,
        f"不支持的文件类型 {suffix or '(无扩展名)'}：本系统只接受 .csv / .xlsx",
    )
